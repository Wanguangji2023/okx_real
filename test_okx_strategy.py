#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
okx_autoTrade_TrytoConnectOKX1.py 策略/工程单元测试（文档 9.3.5 / 10.9）
只测纯逻辑与本地账本读写：用内存假客户端 + 临时文件目录，不触碰任何真实 API 与真实账本。
运行: /home/ubuntu/venv/bin/python3 test_okx_strategy.py
"""
import os
import json
import signal
import socket
import subprocess
import sys
import tempfile
import threading
import time
import types
import unittest
from datetime import datetime, timedelta

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import okx_autoTrade_TrytoConnectOKX1 as mod


class TempFilesMixin:
    """把所有 CSV 持久化路径指到临时目录，测试绝不落在真实数据目录"""

    def setUp(self):
        self.tmpdir = tempfile.mkdtemp(prefix="okx_test_")

        def tmp(name):
            return os.path.join(self.tmpdir, name)

        self._orig_paths = {}
        for attr in ("POSITIONS_FILE", "COOLDOWN_FILE", "HIGH_LOW_FILE", "PAIR_STATUS_FILE",
                     "BUY_QUEUE_FILE", "FUNDS_FILE", "TRADE_FILE", "BLACKLIST_FILE",
                     "UNKNOWN_BLACKLIST_FILE", "OBSERVE_LIST_FILE", "OUT_OF_RANGE_FILE",
                     "CONFIG_SNAPSHOT_FILE"):
            # openpyxl 只认 .xlsx 后缀（读的时候按扩展名拒绝），超范围审计表必须保持 xlsx 后缀
            if attr == "OUT_OF_RANGE_FILE":
                ext = ".xlsx"
            elif attr == "CONFIG_SNAPSHOT_FILE":
                ext = ".json"
            else:
                ext = ".csv"
            self._orig_paths[attr] = getattr(mod, attr)
            setattr(mod, attr, tmp(attr + ext))
        mod._CSV_ROW_CACHE.clear()
        mod._XLSX_SET_CACHE.clear()

    def tearDown(self):
        for attr, val in self._orig_paths.items():
            setattr(mod, attr, val)
        mod._CSV_ROW_CACHE.clear()


def _drop_module_open():
    """还原被替换的 mod.open（_make_read_fail 用模块全局 open 拦截读操作）"""
    if hasattr(mod, "open"):
        del mod.open


def make_state(inst_id="TEST-USDT", **over):
    st = object.__new__(mod.SymbolState)
    st.inst_id = inst_id
    st.logger = mod.Logger(os.devnull)
    st.params = {"buy_ratio": 0.01, "cooldown_days": 7, "total_funds": 1000000}
    st.client = None
    st.position_only = False
    st.min_buy_amount = 100.0
    st.blacklist = set()
    st.unknown_blacklist = set()
    st.observe_list = set()
    st.range_high = 2.0
    st.range_low = 0.5
    st.out_of_range = False
    st._okx_available = True
    st._net_failed = False
    st.positions = {}
    st.cooldown = {}
    st.high_low = {}
    st.pair_status = {}
    st.buy_queue = []
    st.high = 2.0
    st.low = 0.5
    st.levels = mod.LayerLadder(0.5, 2.0)
    st.has_position = False
    st.position_price = 0.0
    st.position_qty = 0.0
    st.position_amount = 0.0
    st.buy_time = ""
    st.buy_reason = ""
    st.pair_type = ""
    st.pair_sell_line = ""
    st.peak_price = 0.0
    st.stop_price = 0.0
    st.profit_triggered = False
    st.check_level = "T3"
    st.last_check_time = ""
    st.max_loss_pct = 0.0
    st.is_paired = False
    st.pair_buy_price = 0.0
    st.pair_buy_time = ""
    st.pair_sell_line_price = 0.0
    st.is_sold = False
    st.sell_time = ""
    st.sell_reason = ""
    st._last_profit = 0.0
    st._last_profit_pct = 0.0
    st._last_processed_price = 0.0
    st._last_processed_ts = 0.0
    st._range_flip_ts = 0.0
    st._balance_insufficient_until = 0.0
    st._buy_blocked_until = 0.0
    st._sell_blocked_until = 0.0
    st._state_dirty = False
    # v5.16/v5.17 新增字段（与 SymbolState.__init__ 保持一致，供各回归测试直接构造状态）
    st.cost_source = ""
    st._dust_check_due_ts = 0.0
    st.last_tick_ts = 0.0
    st._last_stale_poll_ts = 0.0
    st._env_missing = False
    st._region_blocked = False   # v5.19：地域/合规限制（不可交易产品）
    # v5.18（文档 C7）：订单执行线程字段（测试默认 None → 同步路径，与旧行为一致）
    st.executor = None
    st.order_in_flight = None
    st.buy_signal_activated = False
    st.buy_signal_lowest_price = 0.0
    st.buy_signal_buy_level = ""
    st.buy_signal_buy_price = 0.0
    st.buy_signal_attempt_time = 0.0
    # v6（设计单 v1.1）：买点状态机 / 已涨幅度 / 清仓规则字段（与 SymbolState.__init__ 对齐）
    st.entry_type = ""
    st.buy_layer_low = 0.0
    st.buy_layer_high = 0.0
    st.risen_low = 0.0
    st.risen_high = 0.0
    st.risen_low_ts = ""
    st.risen_high_ts = ""
    st.risen_armed = False
    st.v6_next_buy = ""
    st.risen_valid_min = 0.05
    st.ref_low_minutes = 120.0
    st.risen_push_throttle = 20.0
    st.v6_sig_type = ""
    st.v6_sig_ref_low = 0.0
    st.v6_sig_ref_low_ts = 0.0
    st._v6_last_attempt_ts = 0.0
    st._v6_risen_pushed = False
    for k, v in over.items():
        setattr(st, k, v)
    return st


class TestPriceLevels(unittest.TestCase):
    def test_level_ordering(self):
        lv = mod.PriceLevels(high=100.0, low=50.0)
        self.assertGreater(lv.pullback_up, lv.pullback2_down)
        self.assertGreater(lv.pullback2_down, lv.callback2_down)
        self.assertGreater(lv.callback2_down, lv.transition2_down)
        self.assertGreater(lv.transition2_down, lv.limit2_down)
        self.assertGreater(lv.limit2_down, lv.low)
        for name, level in lv.buy_levels.items():
            self.assertLess(level, lv.high)
            self.assertGreater(level, lv.low)

    def test_quarter_anchors(self):
        # R4：四个族群各占 1/4 带，锚点必须是 high / 3/4 / 1/2 / 1/4 / low
        lv = mod.PriceLevels(high=2.0, low=0.5)
        self.assertAlmostEqual(lv.mid, 1.25)
        self.assertAlmostEqual(lv.high_mid, 1.625)     # 旧实现这里是 1.25（与 mid 同值）
        self.assertAlmostEqual(lv.low_mid, 0.875)
        self.assertAlmostEqual(lv.pullback2_down, 1.71875)
        self.assertAlmostEqual(lv.callback2_down, 1.34375)
        # 过渡/极限两族不受影响（锚点本就是 mid 与 low 之间）
        self.assertAlmostEqual(lv.transition2_down, 0.96875)
        self.assertAlmostEqual(lv.limit2_down, 0.59375)

    def test_pairs_have_positive_gross_margin(self):
        # R4 核心回归：回调买点曾与回调卖点同为中点 → 毛利 0（扣双边手续费后必亏）
        lv = mod.PriceLevels(high=2.0, low=0.5)
        for buy_key, (sell_key, name) in lv.pair_map.items():
            buy, sell = lv.buy_levels[buy_key], lv.sell_levels[sell_key]
            self.assertGreater(sell, buy, f"{name}族买卖线相等或倒挂")
            margin = (sell - buy) / buy
            self.assertGreater(margin, 0.05, f"{name}族毛利过薄: {margin:.4f}")
            self.assertGreaterEqual(margin, 2 * mod.TAKER_FEE_RATE * 2,
                                    f"{name}族毛利不足以覆盖进出双边手续费")

    def test_degenerate_range(self):
        lv = mod.PriceLevels(high=1.0, low=1.0)
        self.assertEqual(lv.high_mid, 1.0)
        self.assertEqual(lv.mid, 1.0)

    def test_update_recomputes_from_new_range(self):
        lv = mod.PriceLevels(2.0, 0.5)
        lv.update(4.0, 1.0)          # 高/低被行情刷新后仍要保持四分结构
        self.assertAlmostEqual(lv.high_mid, 3.25)
        self.assertAlmostEqual(lv.callback2_down, 2.6875)
        self.assertGreater(lv.sell_levels["callback_up"], lv.buy_levels["callback2_down"])

    def test_pair_map_complete(self):
        lv = mod.PriceLevels(100.0, 50.0)
        for name, (sell_key, _name2) in lv.pair_map.items():
            self.assertIn(sell_key, lv.sell_levels)
            self.assertGreater(lv.sell_levels[sell_key], lv.buy_levels[name])


class TestStopLossTable(unittest.TestCase):
    """R13：136 行硬表换成线性公式后，必须与 v5.11 原表逐档、档内逐点完全等价"""

    OLD_TABLE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "doc_txt",
                             "profit_stop_map_v5.11.json")

    def _old_table(self):
        import json
        with open(self.OLD_TABLE, encoding="utf-8") as f:
            return [(float(p), float(s)) for p, s in json.load(f)]

    def test_thresholds(self):
        self.assertAlmostEqual(mod.get_stop_loss(0.064), 0.015)   # 未达最低档 → 默认 1.5%
        self.assertAlmostEqual(mod.get_stop_loss(0.065), 0.015)
        self.assertAlmostEqual(mod.get_stop_loss(0.100), 0.0115)
        self.assertAlmostEqual(mod.get_stop_loss(0.15), 0.0065)
        self.assertEqual(mod.get_stop_loss(0.20), 0.0)            # ≥20% → 立即清仓档

    def test_matches_old_table_at_every_bucket(self):
        for p, s in self._old_table():
            self.assertAlmostEqual(mod.get_stop_loss(p), s, places=12,
                                   msg=f"档位 {p} 公式 {mod.get_stop_loss(p)} ≠ 原表 {s}")

    def test_matches_old_table_inside_buckets(self):
        table = self._old_table()
        for i, (p, s) in enumerate(table):
            hi = table[i + 1][0] if i + 1 < len(table) else 0.2
            if p >= 0.2:
                continue
            for k in range(1, 10):
                x = p + (hi - p) * k / 10.0        # 档内取值仍应按"档低边"取表值
                self.assertAlmostEqual(mod.get_stop_loss(x), s, places=12,
                                       msg=f"档内 {x} 公式 {mod.get_stop_loss(x)} ≠ 原表 {s}")

    def test_outside_table_edges(self):
        self.assertAlmostEqual(mod.get_stop_loss(0.0), 0.015)
        self.assertAlmostEqual(mod.get_stop_loss(0.0649), 0.015)
        self.assertEqual(mod.get_stop_loss(0.25), 0.0)
        self.assertEqual(mod.get_stop_loss(0.9), 0.0)

    def test_monotonic(self):
        prev = None
        for p, _s in self._old_table():
            s = mod.get_stop_loss(p)
            if prev is not None:
                self.assertLessEqual(s, prev)      # 盈利越高，回撤容忍越小（永不放宽）
            prev = s


class TestBreakeven(unittest.TestCase):
    def test_exact_formula(self):
        # 成本 1000 USDT 买到 999 个币（买入扣费 1 个），单位成本 = 1000/999，卖出再扣 0.1%
        st = make_state(has_position=True, position_price=1.0, position_qty=999.0,
                        position_amount=1000.0)
        be = st._breakeven_price()
        expect = 1000.0 / 999.0 / (1 - mod.TAKER_FEE_RATE)
        self.assertAlmostEqual(be, expect)
        self.assertGreater(be, 1.0)                       # 保本价必须高于成交均价
        self.assertTrue(st._is_breakeven(be))
        self.assertFalse(st._is_breakeven(be * 0.999))

    def test_fallback_without_amount(self):
        st = make_state(has_position=True, position_price=2.0, position_qty=0.0, position_amount=0.0)
        self.assertAlmostEqual(st._breakeven_price(), 2.0 * (1 + 2 * mod.TAKER_FEE_RATE))

    def test_no_position(self):
        st = make_state()
        self.assertEqual(st._breakeven_price(), 0.0)


class TestCanSell(TempFilesMixin, unittest.TestCase):
    def test_requires_position(self):
        st = make_state()
        ok, msg = st._can_sell(1.0, 0.05)
        self.assertFalse(ok)

    def test_min_profit_gate(self):
        st = make_state(has_position=True, position_price=1.0)
        self.assertFalse(st._can_sell(1.04, 0.05)[0])
        self.assertTrue(st._can_sell(1.05, 0.05)[0])

    def test_none_gate(self):
        # min_profit_pct=None：保本价判定由调用方完成，这里必须放行
        st = make_state(has_position=True, position_price=1.0)
        self.assertTrue(st._can_sell(0.5, None)[0])


class TestCanBuy(TempFilesMixin, unittest.TestCase):
    def test_blocks(self):
        st = make_state(position_only=True)
        self.assertFalse(st._can_buy()[0])
        # v6：超范围剔除已废除（价格越界由换层处理），不再作为买入门槛
        st = make_state(has_position=True)
        self.assertFalse(st._can_buy()[0])
        st = make_state(blacklist={"TEST-USDT"})
        self.assertFalse(st._can_buy()[0])

    def test_local_ledger_duplication_guard(self):
        # 本地账本已有持仓（内存标记未同步）也必须拒绝重复买入
        mod.save_positions({"TEST-USDT": {"position_price": 1, "position_qty": 5,
                                          "position_amount": 5}})
        st = make_state()
        ok, msg = st._can_buy()
        self.assertFalse(ok)
        self.assertTrue(st.has_position)


class TestTakeProfit(TempFilesMixin, unittest.TestCase):
    def _armed_state(self):
        return make_state(
            has_position=True, position_price=1.0, position_qty=999.0, position_amount=1000.0,
            peak_price=1.07, profit_triggered=True, stop_price=1.07 * (1 - 0.014) + 0.0001,
        )

    def test_20pct_clears(self):
        st = self._armed_state()
        st.client = None
        sells = []
        st._execute_sell = lambda price, reason, min_profit_pct=0.05, market=False: (sells.append(reason), "SELL")[1]
        out = st._check_take_profit(1.21)
        self.assertEqual(out, "SELL")
        self.assertIn("≥20%", sells[0])

    def test_activation_at_6_5pct(self):
        st = make_state(has_position=True, position_price=1.0, position_qty=999.0,
                        position_amount=1000.0, peak_price=1.0)
        st._push_notify = lambda t: None
        st._save_state = lambda force=False: None
        st._execute_sell = lambda *a, **k: None
        self.assertFalse(st.profit_triggered)
        st._check_take_profit(1.066)
        self.assertTrue(st.profit_triggered)
        self.assertGreaterEqual(st.stop_price, 1.05)   # 止盈价地板 = 成本×1.05

    def test_breakeven_mixed_semantics(self):
        # 止盈价被跌穿且价格 ≥ 保本价 → 立即市价清仓
        st = self._armed_state()
        be = st._breakeven_price()
        sells = []
        st._execute_sell = lambda price, reason, min_profit_pct=0.05, market=False: (sells.append((price, min_profit_pct)), "SELL")[1]
        price = be * 1.001
        out = st._check_take_profit(price)
        self.assertEqual(out, "SELL")
        self.assertIsNone(sells[0][1])   # 豁免 5% 门槛，走保本价口径

    def test_below_breakeven_keeps_position(self):
        # 止盈价被跌穿但已跌破保本价 → 不亏本卖出，激活状态保留
        st = self._armed_state()
        be = st._breakeven_price()
        st._execute_sell = lambda *a, **k: None
        out = st._check_take_profit(be * 0.99)
        self.assertIsNone(out)
        self.assertTrue(st.profit_triggered)

    def test_stop_price_ratchet(self):
        st = make_state(has_position=True, position_price=1.0, position_qty=999.0,
                        position_amount=1000.0, peak_price=1.0, stop_price=1.10,
                        profit_triggered=True)
        st._push_notify = lambda t: None
        st._save_state = lambda force=False: None
        sold = []
        st._execute_sell = lambda *a, **k: sold.append(a) or None
        st._check_take_profit(1.07)     # 价格回落：峰值更新，止盈价只上不下
        self.assertAlmostEqual(st.peak_price, 1.07)
        self.assertGreaterEqual(st.stop_price, 1.10)
        # 1.07 ≤ 止盈价且 ≥ 保本价 → 应触发离场（本例验证棘轮不阻碍正常离场路径）
        self.assertEqual(len(sold), 1)


class TestExecuteSellCostScaling(TempFilesMixin, unittest.TestCase):
    """10.1：平台量偏离时成本额必须等比缩放，保住单位成本"""

    class FakeClient:
        def __init__(self, avail, cash, frozen, fill):
            self.avail, self.cash, self.frozen, self.fill = avail, cash, frozen, fill

        def get_balance_detail(self, ccy):
            return (self.avail, self.cash, self.frozen)

        def get_instrument(self, inst_id):
            return {"minSz": "0.000001", "lotSz": "0.00000001", "baseCcy": "TEST"}

        def market_sell_base(self, inst_id, qty, px=None):
            f = dict(self.fill)
            f["qty"] = qty
            return f

    def test_amount_scaled_not_flattened(self):
        st = make_state(has_position=True, position_price=1.0, position_qty=1000.0,
                        position_amount=1000.0)
        st.client = self.FakeClient(avail=900.0, cash=900.0, frozen=0.0,
                                    fill={"price": 1.1, "proceeds": 990.0,
                                          "ord_id": "1", "fee": -0.99, "fee_ccy": "USDT",
                                          "state": "filled"})
        st._push_notify = lambda t: None
        st._save_state = lambda force=False: None
        st._add_to_buy_queue = lambda s: None
        st._remove_from_buy_queue = lambda: None
        st._check_take_profit = None
        st.levels = mod.PriceLevels(2.0, 0.5)
        out = st._execute_sell(1.1, "测试卖出", min_profit_pct=None)
        self.assertIsNotNone(out)
        # 缩放后单位成本 = 1000×(900/1000)/900 = 1.0，未被 position_price×avail 拍平
        self.assertAlmostEqual(st.position_amount, 0.0)   # 卖出后清零
        self.assertTrue(st.has_position is False)

    def test_intermediate_unit_cost_preserved(self):
        # 10.1 核心口径：等比缩放保住"每枚到手币的成本"（含买入手续费），
        # 而旧的 position_price×avail 写法会把单位成本拍平回成交均价、低估保本价
        avg_px, amount, qty, avail = 1.0, 1000.0, 999.0, 899.1
        scaled = amount * (avail / qty)
        flattened = avg_px * avail
        self.assertAlmostEqual(scaled / avail, amount / qty)          # 单位成本不变
        self.assertGreater(scaled / avail, flattened / avail)         # 缩放口径高于拍平口径


class TestFrozenRetry(TempFilesMixin, unittest.TestCase):
    """9.1.2：可用 0 但账户有量 = 挂单占用，保留账本并延后重试，绝不清除"""

    class FrozenClient:
        def get_balance_detail(self, ccy):
            return (0.0, 100.0, 100.0)

        def get_instrument(self, inst_id):
            return {"minSz": "1", "lotSz": "0.00000001", "baseCcy": "TEST"}

    def test_keeps_ledger(self):
        st = make_state(has_position=True, position_price=1.0, position_qty=100.0,
                        position_amount=100.0)
        st.client = self.FrozenClient()
        st._push_notify = lambda t: None
        st._save_state = lambda force=False: None
        before = st.position_qty
        out = st._execute_sell(1.1, "测试")
        self.assertIsNone(out)
        self.assertTrue(st.has_position)
        self.assertEqual(st.position_qty, before)
        self.assertEqual(st.position_amount, 100.0)   # 成本基准不得被 avail=0 归零
        self.assertGreater(st._sell_blocked_until, 0.0)

    class EmptyClient(FrozenClient):
        def get_balance_detail(self, ccy):
            return (0.0, 0.0, 0.0)

    def test_true_absent_clears_ledger(self):
        # 总量也为 0 = 账本失真，允许清除
        st = make_state(has_position=True, position_price=1.0, position_qty=100.0,
                        position_amount=100.0)
        st.client = self.EmptyClient()
        st._push_notify = lambda t: None
        st._save_state = lambda force=False: None
        out = st._execute_sell(1.1, "测试")
        self.assertIsNone(out)
        self.assertFalse(st.has_position)
        self.assertEqual(st.position_qty, 0.0)


def make_pool_workbook(path, coins, extra_sheets=None):
    """v5.13 股票池格式：sheet1"层计算"（A=股票代码、C=最新价格区间、D=历史最低、F=历史最高），
    sheet2"层计算幅度"（A=币对、C=同层最低→ref_low、D=同层最高→ref_high）。
    coins: [(inst_id, ref_low, ref_high, band_label)]；extra_sheets: [(表名, [行, ...])]"""
    import openpyxl
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "层计算"
    ws1.append(["股票代码", "最新价格", "最新价格区间", "历史最低价格", "历史最低区间", "历史最高价格"])
    ws2 = wb.create_sheet("层计算幅度")
    ws2.append(["股票代码", "现价", "同层最低", "同层同高"])
    for inst, low, high, label in coins:
        ws1.append([inst, low, label, low, label, high])
        ws2.append([inst, low, low, high])
    for name, rows in (extra_sheets or []):
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(list(r))
    wb.save(path)
    return path


class PoolDirMixin:
    """把 CONFIG_DIR 指到临时股票池目录；单测自行写入 xlsx 后再 load_config"""

    def _patch_pool_dir(self):
        self._orig_pool_dir = mod.CONFIG_DIR
        self._orig_filters = mod.LAYER_FILTERS
        mod.LAYER_FILTERS = []
        self.pool_dir = os.path.join(self.tmpdir, "pool")
        os.makedirs(self.pool_dir, exist_ok=True)
        mod.CONFIG_DIR = self.pool_dir

    def _restore_pool_dir(self):
        mod.CONFIG_DIR = self._orig_pool_dir
        mod.LAYER_FILTERS = self._orig_filters


class TestConfigChineseKeys(TempFilesMixin, PoolDirMixin, unittest.TestCase):
    """10.2：中文键名映射必须生效，未知键忽略并点名（v5.13：股票池来自 INPUT_DATA 目录）"""

    def _isolate_env(self):
        # .env 里的 COOLDOWN_DAYS/BUY_RATIO 会覆盖配置表，单测须隔离环境变量
        saved = {}
        for k in ("COOLDOWN_DAYS", "BUY_RATIO"):
            saved[k] = os.environ.pop(k, None)
        return saved

    def _restore_env(self, saved):
        for k, v in saved.items():
            if v is not None:
                os.environ[k] = v

    def test_aliases(self):
        saved = self._isolate_env()
        self._patch_pool_dir()
        try:
            make_pool_workbook(
                os.path.join(self.pool_dir, "demo.xlsx"),
                [("AAA-USDT", 0.5, 2.0, "8-16")],
                extra_sheets=[("系统参数", [
                    ["参数", "值"],
                    ["单次买入比例", 0.02],
                    ["静默期", 10],
                    ["交易模式", "demo"],
                    ["某个不存在的参数", 1],
                ])])
            config, params = mod.load_config()
        finally:
            self._restore_pool_dir()
            self._restore_env(saved)
        self.assertIn("AAA-USDT", config)
        # 旧 B=新 D（同层最高→ref_high）、旧 C=新 C（同层最低→ref_low）
        self.assertAlmostEqual(config["AAA-USDT"]["ref_high"], 2.0)
        self.assertAlmostEqual(config["AAA-USDT"]["ref_low"], 0.5)
        self.assertAlmostEqual(params["buy_ratio"], 0.02)
        self.assertEqual(params["cooldown_days"], 10)

    def test_env_override(self):
        orig_env = os.environ.get("BUY_RATIO")
        saved = self._isolate_env()
        self._patch_pool_dir()
        os.environ["BUY_RATIO"] = "0.03"
        os.environ.pop("COOLDOWN_DAYS", None)
        try:
            make_pool_workbook(os.path.join(self.pool_dir, "demo.xlsx"), [])
            _config, params = mod.load_config()
        finally:
            self._restore_pool_dir()
            os.environ.pop("COOLDOWN_DAYS", None)
            if orig_env is None:
                os.environ.pop("BUY_RATIO", None)
            else:
                os.environ["BUY_RATIO"] = orig_env
            self._restore_env(saved)
        self.assertAlmostEqual(params["buy_ratio"], 0.03)


class TestConfigSnapshotFallback(TempFilesMixin, PoolDirMixin, unittest.TestCase):
    """v5.12 R16：股票池目录/文件读不出来时的三级降级（xlsx → 快照 → 默认参数 + 持仓管理）
    v5.13：数据源改为 INPUT_DATA 目录内全部 xlsx（sheet2 层计算幅度：C=同层最低、D=同层最高）"""

    def _isolate_env(self):
        saved = {}
        for k in ("COOLDOWN_DAYS", "BUY_RATIO"):
            saved[k] = os.environ.pop(k, None)
        return saved

    def _restore_env(self, saved):
        for k, v in saved.items():
            if v is not None:
                os.environ[k] = v

    def _make_workbook(self, rows):
        # rows: [inst_id, ref_high, ref_low] → 新池 sheet2 列序：C=同层最低、D=同层最高
        make_pool_workbook(self._cfg_path, [(r[0], r[2], r[1], "8-16") for r in rows])

    def setUp(self):
        super().setUp()
        self._env = self._isolate_env()
        self._patch_pool_dir()
        self._cfg_path = os.path.join(self.pool_dir, "config.xlsx")
        self._pushed = []
        self._orig_push = mod.push_notify_throttled
        mod.push_notify_throttled = lambda msg, key=None, title=None: self._pushed.append(key)

    def tearDown(self):
        mod.push_notify_throttled = self._orig_push
        self._restore_pool_dir()
        self._restore_env(self._env)
        super().tearDown()

    def test_success_writes_snapshot(self):
        self._make_workbook([["AAA-USDT", 2.0, 0.5], ["BBB-USDT", 4.0, 1.0]])
        config, params, source = mod.load_config_with_fallback()
        self.assertEqual(source, "xlsx")
        self.assertEqual(len(config), 2)
        self.assertTrue(os.path.exists(mod.CONFIG_SNAPSHOT_FILE))
        with open(mod.CONFIG_SNAPSHOT_FILE, encoding="utf-8") as f:
            snap = json.load(f)
        self.assertEqual(sorted(snap["config"]), ["AAA-USDT", "BBB-USDT"])
        self.assertTrue(snap["saved_at"])

    def test_corrupt_workbook_falls_back_to_snapshot(self):
        self._make_workbook([["AAA-USDT", 2.0, 0.5]])
        mod.load_config_with_fallback()
        with open(self._cfg_path, "wb") as f:
            f.write(b"this is not a zip archive")
        config, params, source = mod.load_config_with_fallback()
        self.assertEqual(source, "snapshot")
        self.assertIn("AAA-USDT", config)
        self.assertAlmostEqual(config["AAA-USDT"]["ref_high"], 2.0)
        # 回退路径也要给全默认参数，MarketEngine 拿到的不能是残缺 dict
        for key in mod.CONFIG_PARAM_DEFAULTS:
            self.assertIn(key, params)

    def test_missing_workbook_and_no_snapshot_returns_defaults(self):
        mod.CONFIG_DIR = os.path.join(self.tmpdir, "never_written_dir")
        config, params, source = mod.load_config_with_fallback()
        self.assertEqual(source, "none")
        self.assertEqual(config, {})
        self.assertEqual(params["trade_mode"], mod.CONFIG_PARAM_DEFAULTS["trade_mode"])
        self.assertEqual(params["cooldown_days"], 7)

    def test_snapshot_rows_are_revalidated(self):
        # 快照同样是手工可编辑的文件：倒挂/非数字的行必须在读回时被剔掉
        os.makedirs(os.path.dirname(mod.CONFIG_SNAPSHOT_FILE), exist_ok=True)
        with open(mod.CONFIG_SNAPSHOT_FILE, "w", encoding="utf-8") as f:
            json.dump({"saved_at": "2026-01-01 00:00:00",
                       "config": {"GOOD-USDT": {"ref_high": 2.0, "ref_low": 0.5},
                                  "INVERTED-USDT": {"ref_high": 0.5, "ref_low": 2.0},
                                  "TEXT-USDT": {"ref_high": "x", "ref_low": 1.0},
                                  "BROKEN-USDT": "不是字典"},
                       "system_params": {"buy_ratio": 0.04, "不认识的键": 1}}, f)
        cfg, params, source = mod.load_config_with_fallback()
        self.assertEqual(source, "snapshot")
        self.assertEqual(sorted(cfg), ["GOOD-USDT"])
        self.assertAlmostEqual(params["buy_ratio"], 0.04)
        self.assertNotIn("不认识的键", params)

    def test_env_override_still_applies_on_fallback(self):
        self._make_workbook([["AAA-USDT", 2.0, 0.5]])
        os.environ["BUY_RATIO"] = "0.05"
        os.environ["COOLDOWN_DAYS"] = "12"
        mod.load_config_with_fallback()
        with open(self._cfg_path, "wb") as f:
            f.write(b"broken")
        _cfg, params, source = mod.load_config_with_fallback()
        self.assertEqual(source, "snapshot")
        self.assertAlmostEqual(params["buy_ratio"], 0.05)
        self.assertEqual(params["cooldown_days"], 12)

    def test_empty_table_does_not_poison_snapshot(self):
        # 用户主动清空股票池（0 个币对）是有效意图：不能因此把上一份好快照抹掉
        self._make_workbook([["AAA-USDT", 2.0, 0.5]])
        mod.load_config_with_fallback()
        with open(mod.CONFIG_SNAPSHOT_FILE, encoding="utf-8") as f:
            before = f.read()
        self._make_workbook([])
        config, _params, source = mod.load_config_with_fallback()
        self.assertEqual(source, "xlsx")
        self.assertEqual(config, {})
        with open(mod.CONFIG_SNAPSHOT_FILE, encoding="utf-8") as f:
            self.assertEqual(f.read(), before)
        # 表格随后坏掉时，仍能救回原来那 1 个币对
        with open(self._cfg_path, "wb") as f:
            f.write(b"broken")
        _cfg, _p, source2 = mod.load_config_with_fallback()
        self.assertEqual(source2, "snapshot")

    def test_corrupt_snapshot_file_is_ignored(self):
        self._make_workbook([["AAA-USDT", 2.0, 0.5]])
        mod.load_config_with_fallback()
        with open(mod.CONFIG_SNAPSHOT_FILE, "w", encoding="utf-8") as f:
            f.write("{ not json")
        with open(self._cfg_path, "wb") as f:
            f.write(b"broken")
        config, _params, source = mod.load_config_with_fallback()
        self.assertEqual(source, "none")
        self.assertEqual(config, {})


class TestCsvCache(TempFilesMixin, unittest.TestCase):
    def test_invalidation_on_write(self):
        mod.save_positions({"A-USDT": {"position_price": 1, "position_qty": 1, "position_amount": 1}})
        rows = mod._read_csv(mod.POSITIONS_FILE, mod.POSITIONS_HEADER)
        self.assertEqual(len(rows), 1)
        # 外部直接改文件（缓存按 stat 指纹失效）
        mod._write_csv_row(mod.POSITIONS_FILE, mod.POSITIONS_HEADER,
                           {"inst_id": "B-USDT", "position_price": 2})
        rows = mod._read_csv(mod.POSITIONS_FILE, mod.POSITIONS_HEADER)
        self.assertEqual(len(rows), 2)

    def test_returned_rows_are_copies(self):
        mod.save_positions({"A-USDT": {"position_price": 1, "position_qty": 1, "position_amount": 1}})
        rows = mod._read_csv(mod.POSITIONS_FILE, mod.POSITIONS_HEADER)
        rows[0]["position_price"] = 999
        again = mod._read_csv(mod.POSITIONS_FILE, mod.POSITIONS_HEADER)
        self.assertEqual(again[0]["position_price"], "1")

    def test_to_float_guards(self):
        self.assertEqual(mod._to_float("abc"), 0.0)
        self.assertEqual(mod._to_float(None, -1), -1.0)
        self.assertEqual(mod._to_float(""), 0.0)
        self.assertFalse(mod._is_number("2026-10-12 15:55:24"))


class TestBuyQueueReadModifyWrite(TempFilesMixin, unittest.TestCase):
    """10.4：队列变更必须基于盘上最新内容，不得用内存陈旧副本覆盖"""

    def test_other_symbol_changes_survive(self):
        mod.save_buy_queue([{"inst_id": "AAA-USDT", "status": "待买入", "ref_high": "2",
                             "ref_low": "0.5", "add_time": "", "last_check_time": ""},
                            {"inst_id": "BBB-USDT", "status": "待买入", "ref_high": "2",
                             "ref_low": "0.5", "add_time": "", "last_check_time": ""}])
        st_a = make_state("AAA-USDT")
        st_b = make_state("BBB-USDT")
        # A 先把 BBB 改成静默期（模拟另一币对的运行期变更），再让 A 改自己那行
        st_b._add_to_buy_queue("静默期_排除")
        st_a._add_to_buy_queue("已持仓_排除")
        rows = {r["inst_id"]: r["status"] for r in mod.load_buy_queue()}
        self.assertEqual(rows["AAA-USDT"], "已持仓_排除")
        self.assertEqual(rows["BBB-USDT"], "静默期_排除")   # 未被陈旧副本打回"待买入"


class TestFlushDirtyStates(TempFilesMixin, unittest.TestCase):
    """v5.10 合并落盘：多个脏币对一次写完，账本内容逐行等价"""

    def test_merge(self):
        engine = mod.MarketEngine(mod.Logger(os.devnull))
        s1 = make_state("A-USDT", has_position=True, position_price=1, position_qty=10,
                        position_amount=10, _state_dirty=True)
        s2 = make_state("B-USDT", has_position=True, position_price=2, position_qty=20,
                        position_amount=40, _state_dirty=True)
        engine.states = {"A-USDT": s1, "B-USDT": s2}
        engine.flush_dirty_states()
        pos = mod.load_positions()
        self.assertIn("A-USDT", pos)
        self.assertIn("B-USDT", pos)
        self.assertFalse(s1._state_dirty)
        self.assertFalse(s2._state_dirty)


class TestWsFrame(TempFilesMixin, unittest.TestCase):
    """10.8：分片帧续接 + 握手残留字节复用（用内存对喂，不走网络）"""

    class FakeSock:
        def __init__(self, chunks):
            self.chunks = list(chunks)

        def recv(self, n):
            if not self.chunks:
                return b""
            return self.chunks.pop(0)

    def _frame(self, payload, opcode=0x1, fin=True):
        first = (0x80 if fin else 0) | opcode
        return bytes([first, len(payload)]) + payload

    def test_fragment_reassembly(self):
        conn = mod.WsConnection(self.FakeSock([self._frame(b'{"a"', fin=False),
                                               self._frame(b':1}', opcode=0x0)]))
        opcode, payload = conn.recv_message()
        self.assertEqual(payload, b'{"a":1}')

    def test_handshake_leftover(self):
        leftover = self._frame(b"pong-left")
        conn = mod.WsConnection(self.FakeSock([]), recv_buf=leftover)
        opcode, payload = conn.recv_message()
        self.assertEqual(payload, b"pong-left")

    def test_control_frame_passthrough(self):
        conn = mod.WsConnection(self.FakeSock([]), recv_buf=self._frame(b"hb", opcode=0x9))
        opcode, payload = conn.recv_message()
        self.assertEqual(opcode, 0x9)
        self.assertEqual(payload, b"hb")

    def test_frame_roundtrip_masking(self):
        payload = "订阅消息".encode("utf-8")
        raw = mod.build_ws_frame(payload)
        # 客户端发出必须带掩码
        self.assertTrue(raw[1] & 0x80)
        # 手工解回
        length = raw[1] & 0x7F
        mask_key = raw[2:6]
        masked = raw[6:6 + length]
        unmasked = bytes(b ^ mask_key[i % 4] for i, b in enumerate(masked))
        self.assertEqual(unmasked, payload)


class TestReadFailureGuard(TempFilesMixin, unittest.TestCase):
    """v5.12 R1：读失败不得当"空表"落盘，否则整表重写会抹掉其他币对的行"""

    def setUp(self):
        super().setUp()
        mod._CSV_READ_FAILED.clear()
        self.addCleanup(mod._CSV_READ_FAILED.clear)

    def _make_read_fail(self, path):
        """只让该文件的读打开失败，写与临时文件不受影响（模拟 EMFILE/OOM 一类瞬时错误）"""
        import builtins

        def fake_open(file, *a, **k):
            mode = a[0] if a else k.get("mode", "r")
            if str(file) == str(path) and "r" in mode and "+" not in mode:
                raise OSError(24, "模拟读失败")
            return builtins.open(file, *a, **k)

        mod.open = fake_open
        self.addCleanup(_drop_module_open)

    def test_failed_read_is_not_cached(self):
        mod.save_positions({"A-USDT": {"position_price": 1, "position_qty": 1, "position_amount": 1}})
        self._make_read_fail(mod.POSITIONS_FILE)
        rows = mod._read_csv(mod.POSITIONS_FILE, mod.POSITIONS_HEADER)
        self.assertEqual(rows, [])
        self.assertNotIn(mod.POSITIONS_FILE, mod._CSV_ROW_CACHE)   # 坏读数绝不进缓存
        self.assertIn(mod.POSITIONS_FILE, mod._CSV_READ_FAILED)

    def test_whole_table_rewrite_refused_after_read_failure(self):
        mod.save_positions({
            "A-USDT": {"position_price": 1, "position_qty": 1, "position_amount": 1},
            "B-USDT": {"position_price": 2, "position_qty": 2, "position_amount": 4},
        })
        # 复现真实路径：读全表失败 → 拿到空内容 → 改一行 → 整表重写（会把 B 抹掉）
        self._make_read_fail(mod.POSITIONS_FILE)
        self.assertEqual(mod.load_positions(), {})
        ok = mod.save_positions({"A-USDT": {"position_price": 1, "position_qty": 1, "position_amount": 1}})
        _drop_module_open()
        self.assertFalse(ok)
        pos = mod.load_positions()
        self.assertIn("B-USDT", pos)
        self.assertIn("A-USDT", pos)   # 整表写被拒，A 的盘上旧行也保持原样

    def test_dirty_flag_kept_when_write_refused(self):
        engine = mod.MarketEngine(mod.Logger(os.devnull))
        s1 = make_state("A-USDT", has_position=True, position_price=1, position_qty=10,
                        position_amount=10, _state_dirty=True)
        engine.states = {"A-USDT": s1}
        self._make_read_fail(mod.POSITIONS_FILE)
        self.assertFalse(engine.flush_dirty_states())
        self.assertTrue(s1._state_dirty)   # 未落盘 → 保持脏，下个周期重试

    def test_pending_high_low_kept_when_write_refused(self):
        mod.save_high_low_entry("A-USDT", 9.0, 1.0)
        self._make_read_fail(mod.HIGH_LOW_FILE)
        self.assertFalse(mod.flush_high_low())
        self.assertIn("A-USDT", mod._HL_PENDING)   # 极值不丢
        self.addCleanup(mod._HL_PENDING.clear)
        mod._CSV_READ_FAILED.discard(mod.HIGH_LOW_FILE)
        _drop_module_open()
        self.assertTrue(mod.flush_high_low())
        self.assertEqual(mod.load_high_low()["A-USDT"]["all_time_high"], 9.0)

    def test_buy_blocked_while_positions_ledger_unreadable(self):
        st = make_state("A-USDT")
        mod._balance_cache.update({"value": 1000.0, "timestamp": time.time()})
        can, _ = st._can_buy()
        self.assertTrue(can)   # 前提：正常情况下确实可买
        mod._CSV_READ_FAILED.add(mod.POSITIONS_FILE)
        can, msg = st._can_buy()
        self.assertFalse(can)
        self.assertIn("读取失败", msg)

    def test_mark_cleared_by_next_good_read(self):
        mod.save_positions({"A-USDT": {"position_price": 1, "position_qty": 1, "position_amount": 1}})
        mod._CSV_READ_FAILED.add(mod.POSITIONS_FILE)
        rows = mod._read_csv(mod.POSITIONS_FILE, mod.POSITIONS_HEADER)
        self.assertEqual(len(rows), 1)
        self.assertNotIn(mod.POSITIONS_FILE, mod._CSV_READ_FAILED)
        self.assertTrue(mod.save_positions({}))


class TestCsvAtomicWrite(TempFilesMixin, unittest.TestCase):
    """v5.12 R2：整表写的临时文件按 pid+线程隔离，且 replace 与写同在一把锁内"""

    def test_tmp_name_differs_per_concurrent_writer(self):
        names = set()
        barrier = threading.Barrier(8)
        def record():
            barrier.wait()   # 让 8 个写者同时存活（ident 会在线程退出后复用）
            names.add(mod._csv_tmp_path(mod.POSITIONS_FILE))
            barrier.wait()
        ts = [threading.Thread(target=record) for _ in range(8)]
        for t in ts:
            t.start()
        for t in ts:
            t.join()
        self.assertEqual(len(names), 8)
        self.assertTrue(all(".tmp." in n for n in names))

    def test_concurrent_whole_table_writes_never_mix(self):
        counts = {5: [f"P{i}-USDT" for i in range(5)], 9: [f"Q{i}-USDT" for i in range(9)]}
        errors = []

        def writer(n):
            try:
                for _ in range(30):
                    rows = {pid: {"position_price": 1, "position_qty": 1, "position_amount": 1}
                            for pid in counts[n]}
                    mod.save_positions(rows)
            except Exception as e:
                errors.append(e)

        ts = [threading.Thread(target=writer, args=(n,)) for n in counts]
        for t in ts:
            t.start()
        for t in ts:
            t.join()
        self.assertEqual(errors, [])
        rows = mod._read_csv(mod.POSITIONS_FILE, mod.POSITIONS_HEADER)
        self.assertIn(len(rows), (5, 9))   # 只会是某个写者的完整表，不会半张表
        names = {r["inst_id"] for r in rows}
        self.assertTrue(all(n.startswith(list(names)[0][0]) for n in names))   # 不混两套前缀
        stray = [f for f in os.listdir(os.path.dirname(mod.POSITIONS_FILE)) if ".tmp." in f]
        self.assertEqual(stray, [])   # 成功后不留临时文件


class TestLedgerLocking(TempFilesMixin, unittest.TestCase):
    """v5.12 R3：后台账户恢复线程的平台对齐必须与行情线程共用一把账本锁"""

    def test_engine_lock_is_shared(self):
        engine = mod.MarketEngine(mod.Logger(os.devnull))
        self.assertIs(engine.lock, mod.ENGINE_LOCK)

    def test_platform_align_waits_for_engine_lock(self):
        mod.save_positions({"A-USDT": {"position_price": 1, "position_qty": 1, "position_amount": 1}})
        done = threading.Event()
        result = {}

        def worker():
            try:
                result["ok"] = mod._apply_platform_positions({})
            except Exception as e:
                result["err"] = e
            finally:
                done.set()

        with mod.ENGINE_LOCK:   # 模拟行情线程正在 on_trade 里读写账本
            t = threading.Thread(target=worker)
            t.start()
            self.assertFalse(done.wait(0.3))   # 锁没释放就不许落盘
        done.wait(5)   # 释放锁后应当立即完成
        t.join(5)
        self.assertNotIn("err", result)
        self.assertTrue(result.get("ok"))


class TestBuyErrorClassification(TempFilesMixin, unittest.TestCase):
    """v5.12 R5：未成交/余额不足是暂时状态，绝不能永久打入未知异常黑名单（那样只能手工改 CSV 恢复）"""

    class BuyClient:
        def __init__(self, result):
            self.result = result

        def get_balance(self, ccy="USDT"):
            return 500.0

        def get_instrument(self, inst_id):
            return {"minSz": "0.0001", "lotSz": "0.0001", "baseCcy": "TEST"}

        def market_buy_quote(self, inst_id, amount, px=None):
            return self.result

    def _try_buy(self, result):
        st = make_state("TEST-USDT")
        st.client = self.BuyClient(result)
        st._push_notify = lambda t: None
        st._save_state = lambda force=False: None
        st._remove_from_buy_queue = lambda: None
        st._add_to_buy_queue = lambda s: None
        orig = mod._okx_client
        mod._okx_client = st.client
        try:
            out = st._execute_buy(1.0, "测试买入", "pullback2_down", "pullback_up", "回撤")
        finally:
            mod._okx_client = orig
        return st, out

    def test_not_filled_cooldowns_instead_of_blacklist(self):
        st, out = self._try_buy({"not_filled": True, "state": "canceled", "ord_id": "1"})
        self.assertIsNone(out)
        self.assertGreater(st._buy_blocked_until, time.time())
        self.assertNotIn("TEST-USDT", mod.load_unknown_blacklist())
        self.assertNotIn("TEST-USDT", mod.load_blacklist())

    def test_empty_reply_cooldowns_instead_of_blacklist(self):
        st, out = self._try_buy(None)
        self.assertIsNone(out)
        self.assertGreater(st._buy_blocked_until, time.time())
        self.assertNotIn("TEST-USDT", mod.load_unknown_blacklist())

    def test_cooldown_blocks_next_buy_attempt(self):
        st, _ = self._try_buy({"not_filled": True, "state": "", "ord_id": ""})
        can, msg = st._can_buy()
        self.assertFalse(can)
        self.assertIn("冷却", msg)

    def test_insufficient_balance_cooldowns(self):
        st, out = self._try_buy({"error": "code=0 msg= data=[{'sCode':'51008','sMsg':'Insufficient USDT balance'}]"})
        self.assertIsNone(out)
        self.assertGreater(st._balance_insufficient_until, time.time())
        self.assertNotIn("TEST-USDT", mod.load_unknown_blacklist())

    def test_compliance_still_blacklisted(self):
        st, out = self._try_buy({"error": "code=51155 msg=Your地区 restriction"})
        self.assertIsNone(out)
        self.assertIn("TEST-USDT", mod.load_blacklist())
        self.assertIn("TEST-USDT", st.blacklist)   # 内存集合同步（R15）

    def test_env_missing_51001_blacklisted(self):
        # v5.19（用户决定）：当前环境无此产品(51001) → 常规黑名单持久化，不进未知异常复检（C11）
        st, out = self._try_buy({"error": "code=51001 msg=Instrument ID, Instrument ID code, or Spread ID doesn't exist."})
        self.assertIsNone(out)
        self.assertIn("TEST-USDT", mod.load_blacklist())
        self.assertIn("TEST-USDT", st.blacklist)                    # 内存集合同步
        self.assertNotIn("TEST-USDT", mod.load_unknown_blacklist()) # 不占 24h 复检通道
        self.assertTrue(st._env_missing)

    def test_unknown_error_still_blacklisted(self):
        st, out = self._try_buy({"error": "code=99999 msg=从来没见过的错误"})
        self.assertIsNone(out)
        self.assertIn("TEST-USDT", mod.load_unknown_blacklist())

    def test_order_limit_goes_to_observe_list(self):
        st, out = self._try_buy({"error": "code=0 msg= data=[{'sCode':'51201','sMsg':'Order exceed market order limit'}]"})
        self.assertIsNone(out)
        self.assertIn("TEST-USDT", mod.load_observe_list())
        self.assertIn("TEST-USDT", st.observe_list)   # 内存集合同步（R15）
        self.assertNotIn("TEST-USDT", mod.load_unknown_blacklist())


class TestSellFailureBackoff(TempFilesMixin, unittest.TestCase):
    """v5.12 R6：卖出失败必须退避，否则每个行情 tick 都在锁里做 10s REST + 下单 + 等成交"""

    class FailingClient:
        def __init__(self):
            self.calls = 0

        def get_balance_detail(self, ccy):
            self.calls += 1
            return (100.0, 100.0, 0.0)

        def get_instrument(self, inst_id):
            return {"minSz": "0.0001", "lotSz": "0.0001", "baseCcy": "TEST"}

        def market_sell_base(self, inst_id, qty, px=None):
            return None   # 下单被拒/未成交

    def test_failure_sets_backoff(self):
        st = make_state(has_position=True, position_price=1.0, position_qty=100.0,
                        position_amount=100.0)
        st.client = self.FailingClient()
        st._push_notify = lambda t: None
        st._save_state = lambda force=False: None
        self.assertIsNone(st._execute_sell(1.1, "止盈清仓", min_profit_pct=None))
        self.assertGreater(st._sell_blocked_until, time.time())
        # 退避期内不再碰平台（连 get_balance_detail 都不该调）
        before = st.client.calls
        self.assertIsNone(st._execute_sell(1.1, "止盈清仓", min_profit_pct=None))
        self.assertEqual(st.client.calls, before)
        self.assertTrue(st.has_position)   # 失败绝不改判持仓


class TestV6BuyFsm(TempFilesMixin, unittest.TestCase):
    """v6（设计单 v1.1）§2：两类买点状态机——激活/撤信号/作废转移/限价买入。
    层梯表 [0.5, 2.0]：回撤下(13/16)=1.71875，高位中间(12/16)=1.625，执行价(12.5/16)=1.671875；
    极限下(1/16)=0.59375，地板=0.5，低位执行价(0.5/16)=0.546875。"""

    def setUp(self):
        super().setUp()
        self.buys = []
        self._orig_throttle = mod.push_notify_throttled
        mod.push_notify_throttled = lambda *a, **k: None

    def tearDown(self):
        mod.push_notify_throttled = self._orig_throttle
        super().tearDown()

    def _state(self, **over):
        st = make_state(**over)
        st._save_state = lambda force=False: None
        st._push_notify = lambda t: None
        st.ref_low_minutes = 0.0     # 测试内免等 120 分钟

        def fake_buy(price, reason, buy_key, sell_key, pair_name, entry_type=""):
            self.buys.append((price, buy_key, sell_key, pair_name, entry_type))
            st._reset_buy_signal()      # 与真实 _buy_apply 一致：成交后复位信号
            return {"qty": 1.0}

        st._execute_buy = fake_buy
        return st

    def _feed(self, st, *prices):
        alerts = []
        for i, p in enumerate(prices):
            alerts = st.update_price(p, 1000.0 + i)
        return alerts

    def test_high_activates_in_zone(self):
        st = self._state()
        self._feed(st, 1.70)                       # 下破 13/16 且 ≥ 12/16 → 高位买点激活
        self.assertEqual(st.v6_sig_type, "high")
        self.assertAlmostEqual(st.v6_sig_ref_low, 1.70)
        self.assertEqual(self.buys, [])

    def test_high_cancels_on_recovery(self):
        st = self._state()
        self._feed(st, 1.70, 1.75)                 # 回升上破 13/16 → 撤信号
        self.assertEqual(st.v6_sig_type, "")
        self.assertEqual(self.buys, [])

    def test_high_void_moves_next_to_low(self):
        st = self._state()
        self._feed(st, 1.70, 1.60)                 # 破 12/16 → 作废，下一买点=低位买点
        self.assertEqual(st.v6_sig_type, "")
        self.assertEqual(st.v6_next_buy, "low")
        self.assertEqual(self.buys, [])

    def test_void_then_low_transfer_executes_low_buy(self):
        # （2026-10-19 钉扎）作废后"下一买点=低位买点"不是死状态：状态机继续执行——
        # 高位被"low"屏蔽不复活；深跌 ≤1/16 激活低信号 → ≤0.5/16 执行价限价买入
        st = self._state()
        self._feed(st, 1.70)                       # 高位买点激活
        self.assertEqual(st.v6_sig_type, "high")
        self._feed(st, 1.60)                       # 破 12/16 → 作废转移
        self.assertEqual(st.v6_next_buy, "low")
        self._feed(st, 1.70)                       # 回升高位带内：高位不复活（被"low"屏蔽）
        self.assertEqual(st.v6_sig_type, "")
        self.assertEqual(self.buys, [])
        self._feed(st, 0.59)                       # 深跌 ≤ 1/16 → 低位买点激活（"low"模式放行）
        self.assertEqual(st.v6_sig_type, "low")
        self._feed(st, 0.54)                       # 静默成立且 ≤ 0.546875 → 限价买入
        self.assertEqual(len(self.buys), 1)
        price, buy_key, sell_key, pair_name, entry_type = self.buys[0]
        self.assertEqual(buy_key, "v6_low")
        self.assertEqual(entry_type, "low")

    def test_flat_next_buy_persists_and_survives_restart(self):
        # （2026-10-19 巡检修复钉扎）作废转移标记写入 buy_queue 行；重启读回，不再复活高位买点
        st = self._state(inst_id="V6Q-USDT")
        self._feed(st, 1.70, 1.60)                 # 激活 → 破 12/16 作废
        self.assertEqual(st.v6_next_buy, "low")
        rows = mod.load_buy_queue()
        row = next(r for r in rows if r.get("inst_id") == "V6Q-USDT")
        self.assertEqual((row.get("v6_next_buy") or ""), "low")
        # 重启：新实例读回标记（层 [0.5, 2.0] 由配置重建）
        st2 = mod.SymbolState("V6Q-USDT", 2.0, 0.5, mod.Logger(os.devnull),
                              {"buy_ratio": 0.01, "cooldown_days": 7})
        self.assertEqual(st2.v6_next_buy, "low")
        st2._save_state = lambda force=False: None
        st2._push_notify = lambda t: None
        st2.blacklist = set(); st2.unknown_blacklist = set(); st2.observe_list = set()
        st2._okx_available = True
        st2.update_price(1.70, 2000.0)             # 高位带内：被 "low" 屏蔽，不复活
        self.assertEqual(st2.v6_sig_type, "")

    def test_flat_layer_shift_marker_persists(self):
        # 上破换层（无持仓）→ 标记 "next_layer" 落 buy_queue 行；重启读回
        st = self._state(inst_id="V6S-USDT")
        self._feed(st, 2.1)                        # 上破天花板 → 换层向上（忽略，不高位接盘）
        self.assertEqual(st.v6_next_buy, "next_layer")
        rows = mod.load_buy_queue()
        row = next(r for r in rows if r.get("inst_id") == "V6S-USDT")
        self.assertEqual((row.get("v6_next_buy") or ""), "next_layer")
        st2 = mod.SymbolState("V6S-USDT", 4.0, 1.0, mod.Logger(os.devnull),
                              {"buy_ratio": 0.01, "cooldown_days": 7})
        self.assertEqual(st2.v6_next_buy, "next_layer")

    def test_high_fires_limit_buy_at_exec(self):
        st = self._state()
        self._feed(st, 1.70, 1.66)                 # ≤ 执行价 1.671875 且静默期已过 → 限价买入
        self.assertEqual(len(self.buys), 1)
        price, buy_key, sell_key, pair_name, entry_type = self.buys[0]
        self.assertEqual(buy_key, "v6_high")
        self.assertEqual(entry_type, "high")
        self.assertLessEqual(price, 1.671875)

    def test_high_blocked_when_next_buy_low(self):
        st = self._state(v6_next_buy="low")
        self._feed(st, 1.70)
        self.assertEqual(st.v6_sig_type, "")
        self.assertEqual(self.buys, [])

    def test_low_activates_and_fires(self):
        st = self._state()
        self._feed(st, 0.59, 0.54)                 # ≤ 1/16 激活 → ≤ 0.546875 执行价买入
        self.assertEqual(len(self.buys), 1)
        price, buy_key, sell_key, pair_name, entry_type = self.buys[0]
        self.assertEqual(buy_key, "v6_low")
        self.assertEqual(entry_type, "low")

    def test_low_blocked_in_next_layer_mode(self):
        st = self._state(v6_next_buy="next_layer")
        self._feed(st, 0.59)
        self.assertEqual(st.v6_sig_type, "")
        self.assertEqual(self.buys, [])

    def test_deep_zone_sets_next_low(self):
        st = self._state()
        self._feed(st, 1.50)                       # < 12/16 且无信号 → 下一买点=低位买点
        self.assertEqual(st.v6_next_buy, "low")
        self.assertEqual(st.v6_sig_type, "")
        rows = mod.load_buy_queue()                # 巡检修复：深区转移同样写入 buy_queue 行
        row = next(r for r in rows if r.get("inst_id") == "TEST-USDT")
        self.assertEqual((row.get("v6_next_buy") or ""), "low")

    def test_exact_floor_activates_and_fires(self):
        # （修复②）恰好等于地板=未破地板（换层判定为严格 <），按设计"不破地板+≤0.5/16 执行价"仍合格
        st = self._state()
        st.update_price(0.5, 1001.0)               # 恰好地板 → 激活（旧行为：不激活）
        self.assertEqual(st.v6_sig_type, "low")
        self.assertEqual(self.buys, [])
        st.update_price(0.5, 1002.0)               # 静默成立且 ≤ 执行价 → 限价买入
        self.assertEqual(len(self.buys), 1)
        price, buy_key, sell_key, pair_name, entry_type = self.buys[0]
        self.assertEqual(buy_key, "v6_low")
        self.assertEqual(entry_type, "low")


class TestV6ClearNextBuy(TempFilesMixin, unittest.TestCase):
    """v6 清仓后下一买点（修复①）：清仓3 破买入层地板当拍换层已完成（层已下移），
    下一买点=下一层高位买点 → 状态机应在（已下移的）当前层按梯表评估，不再多屏蔽一层。"""

    class SellClient:
        def __init__(self, qty):
            self.qty = qty

        def get_balance(self, ccy="USDT"):
            return 100000.0

        def get_balance_detail(self, ccy):
            return (self.qty, self.qty, 0.0)

        def get_instrument(self, inst_id):
            return {"minSz": "0.000001", "lotSz": "0.00000001",
                    "baseCcy": "TEST", "tickSz": "0.0001"}

        def market_sell_base(self, inst_id, qty, px=None):
            p = px or 1.0
            return {"price": p, "qty": qty, "proceeds": p * qty, "ord_id": "s1",
                    "fee": 0.0, "fee_ccy": "USDT", "state": "filled"}

    def setUp(self):
        super().setUp()
        self._orig_throttle = mod.push_notify_throttled
        mod.push_notify_throttled = lambda *a, **k: None
        self._orig_client = mod._okx_client
        mod._okx_client = self.SellClient(2000.0)

    def tearDown(self):
        mod.push_notify_throttled = self._orig_throttle
        mod._okx_client = self._orig_client
        super().tearDown()

    def test_qing3_next_buy_evaluates_in_shifted_layer(self):
        st = make_state("Q3-USDT", has_position=True, position_price=0.55,
                        position_qty=2000.0, position_amount=1100.0,
                        entry_type="low", buy_layer_low=0.5, buy_layer_high=2.0,
                        risen_low=0.55, risen_high=0.55)
        st.client = mod._okx_client
        st._push_notify = lambda t: None
        st._save_state = lambda force=False: None
        buys = []

        def fake_buy(price, reason, buy_key, sell_key, pair_name, entry_type=""):
            buys.append((price, buy_key, entry_type))
            st._reset_buy_signal()
            return {"qty": 1.0}

        st._execute_buy = fake_buy
        st.update_price(0.49, 1001.0)              # 破买入层地板（0.49<0.5，未达 12.5% 止损线 0.48125）→ 清仓3，层已下移至 [0.25, 1.0]
        self.assertFalse(st.has_position)
        self.assertEqual(st.v6_next_buy, "")       # 修复①：保持空标记（旧行为="next_layer" 多屏蔽一层）
        self.assertAlmostEqual(st.levels.low, 0.25)
        self.assertAlmostEqual(st.levels.high, 1.0)
        # v6 起卖出不再写静默期（本批废止）：复买不再被冻结
        self.assertNotIn("Q3-USDT", st.cooldown)
        self.assertFalse(st._is_in_cooldown())
        st.ref_low_minutes = 0.0
        st.update_price(0.85, 1002.0)              # < 回撤下 0.859375 且 ≥ 高位中间 0.8125 → 激活
        self.assertEqual(st.v6_sig_type, "high")
        self.assertEqual(buys, [])
        st.update_price(0.83, 1003.0)              # ≤ 执行价 0.8359375 → v6_high 限价买入
        self.assertEqual(len(buys), 1)
        self.assertEqual(buys[0][1], "v6_high")
        self.assertEqual(buys[0][2], "high")

    def test_qing2_marker_persists_to_buy_queue(self):
        # （2026-10-19 巡检修复钉扎）清仓2（高位买入破买入层高位中间）→ 下一买点="low"，
        # 卖出成交后写入 buy_queue 行；重启读回（此前仅存内存，重启丢失）
        st = make_state("Q2-USDT", has_position=True, position_price=1.60,
                        position_qty=2000.0, position_amount=3200.0,
                        entry_type="high", buy_layer_low=0.5, buy_layer_high=2.0,
                        risen_low=1.60, risen_high=1.60)
        st.client = mod._okx_client
        st._push_notify = lambda t: None
        st._save_state = lambda force=False: None
        st.update_price(1.50, 1001.0)              # 破买入层 12/16=1.625 → 清仓2
        self.assertFalse(st.has_position)
        self.assertEqual(st.v6_next_buy, "low")
        rows = mod.load_buy_queue()
        row = next(r for r in rows if r.get("inst_id") == "Q2-USDT")
        self.assertEqual((row.get("v6_next_buy") or ""), "low")
        st2 = mod.SymbolState("Q2-USDT", 2.0, 0.5, mod.Logger(os.devnull),
                              {"buy_ratio": 0.01, "cooldown_days": 7})
        self.assertEqual(st2.v6_next_buy, "low")


class TestV6StopLoss(TempFilesMixin, unittest.TestCase):
    """铁律（2026-10-19 用户授权）：个股最大亏损 12.5% 底线——触发即市价止损清仓；
    下一买点按买入类型（同回本清仓 E4）；优先于其余清仓规则。

    以卖出调用里的 px 是否为 None 区分：止损=市价（px=None），其余清仓=限价@最新价（px=价格）。
    """

    class SellClient:
        def __init__(self, qty, market_px):
            self.qty = qty
            self.market_px = market_px
            self.calls = []

        def get_balance(self, ccy="USDT"):
            return 100000.0

        def get_balance_detail(self, ccy):
            return (self.qty, self.qty, 0.0)

        def get_instrument(self, inst_id):
            return {"minSz": "0.000001", "lotSz": "0.00000001",
                    "baseCcy": "TEST", "tickSz": "0.0001"}

        def market_sell_base(self, inst_id, qty, px=None):
            self.calls.append({"inst_id": inst_id, "qty": qty, "px": px})
            p = self.market_px if px is None else px   # px=None → 市价单，按当时市价成交
            return {"price": p, "qty": qty, "proceeds": p * qty, "ord_id": "s1",
                    "fee": 0.0, "fee_ccy": "USDT", "state": "filled"}

    def setUp(self):
        super().setUp()
        self._orig_throttle = mod.push_notify_throttled
        mod.push_notify_throttled = lambda *a, **k: None
        self._orig_client = mod._okx_client

    def tearDown(self):
        mod.push_notify_throttled = self._orig_throttle
        mod._okx_client = self._orig_client
        super().tearDown()

    def _mk(self, inst, market_px, **over):
        client = self.SellClient(over.get("position_qty", 2000.0), market_px)
        mod._okx_client = client
        st = make_state(inst, has_position=True, **over)
        st.client = client
        st._push_notify = lambda t: None
        st._save_state = lambda force=False: None
        return st, client

    def test_stop_loss_market_sell_and_marker_persist(self):
        st, client = self._mk("SL1-USDT", 0.874, position_price=1.0, position_qty=2000.0,
                              position_amount=2000.0, entry_type="high",
                              buy_layer_low=0.5, buy_layer_high=2.0,
                              risen_low=1.0, risen_high=1.0, risen_armed=True)
        st.update_price(0.874, 1001.0)                 # 亏损 -12.6%，触发 12.5% 底线
        self.assertFalse(st.has_position)
        self.assertEqual(len(client.calls), 1)
        self.assertIsNone(client.calls[0]["px"])       # 市价单（限价清仓会带 px=最新价）
        self.assertEqual(st.v6_next_buy, "low")        # 高位买入 → 本层低位买点
        rows = mod.load_buy_queue()
        row = next(r for r in rows if r.get("inst_id") == "SL1-USDT")
        self.assertEqual((row.get("v6_next_buy") or ""), "low")
        st2 = mod.SymbolState("SL1-USDT", 2.0, 0.5, mod.Logger(os.devnull),
                              {"buy_ratio": 0.01, "cooldown_days": 7})
        self.assertEqual(st2.v6_next_buy, "low")       # 重启读回

    def test_stop_loss_boundary_exact_line_fires_and_above_not(self):
        st, client = self._mk("SL2-USDT", 0.876, position_price=1.0, position_qty=2000.0,
                              position_amount=2000.0, entry_type="legacy",
                              risen_low=1.0, risen_high=1.0)
        st.update_price(0.876, 1001.0)                 # -12.4%：未到线，不清仓
        self.assertTrue(st.has_position)
        self.assertEqual(client.calls, [])
        st.update_price(0.875, 1002.0)                 # -12.5%：恰达底线 → 市价止损
        self.assertFalse(st.has_position)
        self.assertEqual(len(client.calls), 1)
        self.assertIsNone(client.calls[0]["px"])
        self.assertEqual(st.v6_next_buy, "")           # legacy → 空（本层高位买点）

    def test_stop_loss_precedes_risen_clear(self):
        # 若没有止损，本例（已涨幅度有效 + armed + 破回撤上）会被清仓1 以限价卖出并置空标记；
        # 实际出现市价（px=None）+ "low" 标记 → 证明止损分支优先执行。
        st, client = self._mk("SL3-USDT", 0.874, position_price=1.0, position_qty=2000.0,
                              position_amount=2000.0, entry_type="high",
                              buy_layer_low=0.5, buy_layer_high=2.0,
                              risen_low=0.90, risen_high=1.0, risen_armed=True)
        st.update_price(0.874, 1001.0)
        self.assertFalse(st.has_position)
        self.assertEqual(len(client.calls), 1)
        self.assertIsNone(client.calls[0]["px"])
        self.assertEqual(st.v6_next_buy, "low")

    def test_stop_loss_next_buy_for_low_entry(self):
        st, client = self._mk("SL4-USDT", 0.48, position_price=0.55, position_qty=2000.0,
                              position_amount=1100.0, entry_type="low",
                              buy_layer_low=0.5, buy_layer_high=2.0,
                              risen_low=0.55, risen_high=0.55)
        st.update_price(0.48, 1001.0)                  # (0.48-0.55)/0.55 = -12.7%，同时已破地板 0.5
        self.assertFalse(st.has_position)
        self.assertIsNone(client.calls[0]["px"])       # 市价止损优先于清仓3（限价）
        self.assertEqual(st.v6_next_buy, "next_layer") # 低位买入 → 下一层高位买点
        rows = mod.load_buy_queue()
        row = next(r for r in rows if r.get("inst_id") == "SL4-USDT")
        self.assertEqual((row.get("v6_next_buy") or ""), "next_layer")


class TestEnvTuning(TempFilesMixin, unittest.TestCase):
    """参数外置化（.env）调参验证 —— 2026-10-19 用户要求"明确要求的参数写入 .env 并可调参"。

    ① 子进程验证：模块级常量（铁律 + 动态止盈口径）在"无 .env + 无环境变量"时回退默认值，
       注入环境变量后全部生效（cwd 选全新临时目录，模拟干净环境）；
    ② 进程内验证：REF_LOW_MINUTES / RISEN_RANGE_MIN_PCT / RISEN_PUSH_THROTTLE_SEC（E3 明确
       要求 .env 可调）由 SymbolState 构造时即时读取；已涨幅度推送实际使用节流值；
    ③ 行为级验证：HARD_STOP_LOSS_PCT / BREAKEVEN_CLEAR_LOSS_PCT / TAKE_PROFIT_* 改动后触发点
       随动（均含默认值对照，证明差异由参数而非其他规则造成）。
    """

    ENV_KEYS = ("REF_LOW_MINUTES", "RISEN_RANGE_MIN_PCT", "RISEN_PUSH_THROTTLE_SEC", "SELL_REPRICE_SEC",
                "HARD_STOP_LOSS_PCT", "BREAKEVEN_CLEAR_LOSS_PCT", "TAKE_PROFIT_ACTIVATE_PCT",
                "TAKE_PROFIT_CLEAR_PCT", "TAKE_PROFIT_TOLERANCE_PCT", "TAKE_PROFIT_TOLERANCE_SLOPE",
                "TAKE_PROFIT_BUCKET", "TAKE_PROFIT_MIN_MARGIN_PCT")

    _PROBE_CODE = (
        "import sys, json; sys.path.insert(0, '/home/ubuntu/okx_real'); "
        "import okx_autoTrade_TrytoConnectOKX1 as m; "
        "print('RESULT:' + json.dumps([m.HARD_STOP_LOSS_PCT, m.BREAKEVEN_CLEAR_LOSS_PCT, "
        "m.STOP_LOSS_TRIGGER_PCT, m.STOP_LOSS_ZERO_PCT, m.STOP_LOSS_BASE, m.STOP_LOSS_SLOPE, "
        "m.STOP_LOSS_BUCKET, m.TAKE_PROFIT_MIN_MARGIN_PCT]))"
    )

    class SellClient:
        def __init__(self, qty, market_px):
            self.qty = qty
            self.market_px = market_px
            self.calls = []

        def get_balance(self, ccy="USDT"):
            return 100000.0

        def get_balance_detail(self, ccy):
            return (self.qty, self.qty, 0.0)

        def get_instrument(self, inst_id):
            return {"minSz": "0.000001", "lotSz": "0.00000001",
                    "baseCcy": "TEST", "tickSz": "0.0001"}

        def market_sell_base(self, inst_id, qty, px=None):
            self.calls.append({"inst_id": inst_id, "qty": qty, "px": px})
            p = self.market_px if px is None else px   # px=None → 市价单，按当时市价成交
            return {"price": p, "qty": qty, "proceeds": p * qty, "ord_id": "s1",
                    "fee": 0.0, "fee_ccy": "USDT", "state": "filled"}

    def setUp(self):
        super().setUp()
        self._orig_throttle = mod.push_notify_throttled
        mod.push_notify_throttled = lambda *a, **k: None
        self._orig_client = mod._okx_client
        self._orig_consts = {k: getattr(mod, k) for k in (
            "HARD_STOP_LOSS_PCT", "BREAKEVEN_CLEAR_LOSS_PCT", "STOP_LOSS_BASE",
            "STOP_LOSS_TRIGGER_PCT", "STOP_LOSS_ZERO_PCT", "STOP_LOSS_SLOPE",
            "STOP_LOSS_BUCKET", "TAKE_PROFIT_MIN_MARGIN_PCT")}

    def tearDown(self):
        for k, v in self._orig_consts.items():
            setattr(mod, k, v)
        mod.push_notify_throttled = self._orig_throttle
        mod._okx_client = self._orig_client
        super().tearDown()

    def _mk(self, inst, market_px, **over):
        client = self.SellClient(over.get("position_qty", 2000.0), market_px)
        mod._okx_client = client
        st = make_state(inst, has_position=True, **over)
        st.client = client
        st._push_notify = lambda t: None
        st._save_state = lambda force=False: None
        return st, client

    # ---------- ① 子进程：模块级常量映射（无 .env → 默认；注入 env → 生效） ----------

    def _probe(self, extra_env):
        env = {k: v for k, v in os.environ.items() if k not in self.ENV_KEYS}
        env.update(extra_env)
        out = subprocess.run([sys.executable, "-c", self._PROBE_CODE], capture_output=True,
                             text=True, env=env, cwd=tempfile.mkdtemp(), timeout=120)
        self.assertEqual(out.returncode, 0, f"probe 进程失败: {out.stderr[-800:]}")
        line = next((l for l in out.stdout.splitlines() if l.startswith("RESULT:")), None)
        self.assertIsNotNone(line, f"probe 未输出 RESULT: {out.stdout[-400:]}")
        return json.loads(line[len("RESULT:"):])

    def test_module_constants_defaults_without_env(self):
        vals = self._probe({})
        self.assertEqual(vals, [0.125, 0.05, 0.065, 0.20, 0.015, 0.1, 0.001, 0.05])

    def test_module_constants_override_from_env(self):
        vals = self._probe({"HARD_STOP_LOSS_PCT": "0.10", "BREAKEVEN_CLEAR_LOSS_PCT": "0.03",
                            "TAKE_PROFIT_ACTIVATE_PCT": "0.05", "TAKE_PROFIT_CLEAR_PCT": "0.15",
                            "TAKE_PROFIT_TOLERANCE_PCT": "0.02", "TAKE_PROFIT_TOLERANCE_SLOPE": "0.2",
                            "TAKE_PROFIT_BUCKET": "0.002", "TAKE_PROFIT_MIN_MARGIN_PCT": "0.02"})
        self.assertEqual(vals, [0.10, 0.03, 0.05, 0.15, 0.02, 0.2, 0.002, 0.02])

    def test_env_invalid_value_falls_back_with_default(self):
        vals = self._probe({"HARD_STOP_LOSS_PCT": "abc"})
        self.assertEqual(vals[0], 0.125)               # 非法值 → 回退默认（env_num 保护）

    # ---------- ② 进程内：SymbolState 构造即时读取（E3 三键 + 节流接线） ----------

    def test_state_params_read_from_env(self):
        overrides = {"REF_LOW_MINUTES": "45", "RISEN_RANGE_MIN_PCT": "0.08", "RISEN_PUSH_THROTTLE_SEC": "5"}
        saved = {k: os.environ.get(k) for k in overrides}
        try:
            os.environ.update(overrides)
            mod._okx_client = self.SellClient(2000.0, 1.0)
            st = mod.SymbolState("ENVT-USDT", 2.0, 0.5, mod.Logger(os.devnull),
                                 {"buy_ratio": 0.01, "cooldown_days": 7})
            self.assertEqual(st.ref_low_minutes, 45.0)
            self.assertEqual(st.risen_valid_min, 0.08)
            self.assertEqual(st.risen_push_throttle, 5.0)
        finally:
            for k, v in saved.items():
                if v is None:
                    os.environ.pop(k, None)
                else:
                    os.environ[k] = v
        st2 = mod.SymbolState("ENVT2-USDT", 2.0, 0.5, mod.Logger(os.devnull),
                              {"buy_ratio": 0.01, "cooldown_days": 7})
        self.assertEqual(st2.ref_low_minutes, 120.0)   # 缺省 → 默认 120 分钟
        self.assertEqual(st2.risen_valid_min, 0.05)    # 缺省 → 默认 5%
        self.assertEqual(st2.risen_push_throttle, 20.0)

    def test_risen_push_throttle_wired(self):
        captured = []
        mod.push_notify_throttled = (lambda text, key, title="", msg_type="text", interval=None:
                                     captured.append(interval))
        st, _ = self._mk("PT-USDT", 0.99, position_price=0.9, position_qty=1000.0,
                         position_amount=900.0, entry_type="high",
                         buy_layer_low=0.5, buy_layer_high=1.0,
                         risen_low=0.5, risen_high=1.0, risen_armed=True)
        st.risen_push_throttle = 7.5                    # 属性值应原样传给推送节流
        st.update_price(0.99, 1001.0)
        self.assertEqual(captured, [7.5])

    # ---------- ③ 行为级：调参后触发点随动（含默认值对照） ----------

    def test_hard_stop_loss_pct_tunable(self):
        base = dict(position_price=1.0, position_qty=2000.0, position_amount=2000.0,
                    entry_type="legacy", risen_low=1.0, risen_high=1.0)
        # 调后 10%：-10.2% 触发市价止损
        mod.HARD_STOP_LOSS_PCT = 0.10
        st, client = self._mk("SLT1-USDT", 0.898, **base)
        st.update_price(0.898, 1001.0)
        self.assertFalse(st.has_position)
        self.assertEqual(len(client.calls), 1)
        self.assertIsNone(client.calls[0]["px"])       # 市价（限价清仓会带 px）
        # 调后 10%：-9.8% 不动
        st2, client2 = self._mk("SLT2-USDT", 0.902, **base)
        st2.update_price(0.902, 1001.0)
        self.assertTrue(st2.has_position)
        self.assertEqual(client2.calls, [])
        # 默认 12.5% 对照：同一价格（-10.2%）不触发止损
        mod.HARD_STOP_LOSS_PCT = self._orig_consts["HARD_STOP_LOSS_PCT"]
        st3, client3 = self._mk("SLT3-USDT", 0.898, **base)
        st3.update_price(0.898, 1001.0)
        self.assertTrue(st3.has_position)
        self.assertEqual(client3.calls, [])

    def test_breakeven_clear_threshold_tunable(self):
        base = dict(position_price=1.0, position_qty=2000.0, position_amount=2000.0,
                    entry_type="low", buy_layer_low=0.5, buy_layer_high=2.0,
                    risen_low=1.0, risen_high=1.0)
        # 默认 5%：-3.1% 记录不满足"深亏"，回到保本价不动
        st, client = self._mk("BET1-USDT", 1.003, **base)
        st.update_price(0.969, 1001.0)                 # 创 -3.1% 记录
        self.assertTrue(st.has_position)
        st.update_price(1.003, 1002.0)                 # 保本价上方：默认门槛不触发
        self.assertTrue(st.has_position)
        self.assertEqual(client.calls, [])
        # 调后 3%：同样 -3.1% 记录 → 回本即清（限价 @ 最新价，非市价）
        mod.BREAKEVEN_CLEAR_LOSS_PCT = 0.03
        st2, client2 = self._mk("BET2-USDT", 1.003, **base)
        st2.update_price(0.969, 1001.0)
        self.assertTrue(st2.has_position)
        st2.update_price(1.003, 1002.0)
        self.assertFalse(st2.has_position)
        self.assertEqual(len(client2.calls), 1)
        self.assertEqual(client2.calls[0]["px"], 1.003)    # 限价单
        self.assertEqual(st2.v6_next_buy, "next_layer")    # 低位买入 → 下一层高位买点

    def test_take_profit_activate_pct_tunable(self):
        def act_state():
            s = make_state(has_position=True, position_price=1.0, position_qty=999.0,
                           position_amount=1000.0, peak_price=1.0)
            s._push_notify = lambda t: None
            s._save_state = lambda force=False: None
            s._execute_sell = lambda *a, **k: None
            return s
        st = act_state()
        st._check_take_profit(1.052)                   # 5.2% < 默认 6.5%：不激活
        self.assertFalse(st.profit_triggered)
        mod.STOP_LOSS_TRIGGER_PCT = 0.05
        st2 = act_state()
        st2._check_take_profit(1.052)                  # 调后 5%：5.2% 即激活
        self.assertTrue(st2.profit_triggered)
        self.assertGreaterEqual(st2.stop_price, 1.05)  # 地板仍生效（默认 5% 保底）

    def test_take_profit_floor_and_clear_pct_tunable(self):
        def act_state(price):
            s = make_state(has_position=True, position_price=1.0, position_qty=999.0,
                           position_amount=1000.0, peak_price=1.0)
            s._push_notify = lambda t: None
            s._save_state = lambda force=False: None
            s._execute_sell = lambda *a, **k: None
            s._check_take_profit(price)
            return s
        st = act_state(1.065)                          # 恰好 6.5% 激活：峰值回撤价 1.04902 < 地板 1.05
        self.assertAlmostEqual(st.stop_price, 1.05, places=6)
        mod.TAKE_PROFIT_MIN_MARGIN_PCT = 0.02          # 地板调到 2%：随峰值回撤价
        st2 = act_state(1.065)
        self.assertAlmostEqual(st2.stop_price, 1.065 * (1 - mod.get_stop_loss(0.065)), places=6)
        self.assertLess(st2.stop_price, 1.05)
        # 立即清仓档（默认 20%）：16% 不清仓；调到 15% 后 16% 立即清仓
        def armed_state():
            s = make_state(has_position=True, position_price=1.0, position_qty=999.0,
                           position_amount=1000.0, peak_price=1.16, profit_triggered=True,
                           stop_price=1.0)
            s._push_notify = lambda t: None
            s._save_state = lambda force=False: None
            s._execute_sell = lambda price, reason, min_profit_pct=0.05, market=False: (
                sells.append(reason), "SELL")[1]
            return s
        sells = []
        st3 = armed_state()
        self.assertIsNone(st3._check_take_profit(1.16))
        self.assertEqual(sells, [])
        mod.STOP_LOSS_ZERO_PCT = 0.15
        st4 = armed_state()
        self.assertEqual(st4._check_take_profit(1.16), "SELL")
        self.assertIn("≥15%", sells[-1])


class TestSignalAndWatchdog(unittest.TestCase):
    """R8：SIGTERM 也要走正常收尾；行情线程静默死亡要被发现并拉起"""

    def setUp(self):
        self._sigs = (signal.SIGTERM, signal.SIGINT, signal.SIGQUIT)
        self._orig_handlers = {s: signal.getsignal(s) for s in self._sigs}
        self._orig_stream = mod.stream_okx
        self._orig_throttle = mod.push_notify_throttled
        self._orig_error = mod.WS_HEALTH.get("last_error")
        self._orig_set = mod.STOP_EVENT.is_set()
        mod.STOP_EVENT.clear()
        mod.push_notify_throttled = lambda *a, **k: None

    def tearDown(self):
        for s, h in self._orig_handlers.items():
            signal.signal(s, h)
        mod.stream_okx = self._orig_stream
        mod.push_notify_throttled = self._orig_throttle
        mod.WS_HEALTH["last_error"] = self._orig_error
        mod.STOP_EVENT.clear()
        if self._orig_set:
            mod.STOP_EVENT.set()

    def _dead_thread(self):
        t = threading.Thread(target=lambda: None)
        t.start()
        t.join()
        return t

    def test_handlers_registered_and_stop_event_set(self):
        mod.install_signal_handlers(mod.Logger(os.devnull))
        for s in self._sigs:
            self.assertIs(signal.getsignal(s), mod._handle_stop_signal)
        mod._handle_stop_signal(signal.SIGTERM, None)
        self.assertTrue(mod.STOP_EVENT.is_set())

    def test_second_stop_signal_forces_exit(self):
        orig_term = signal.getsignal(signal.SIGTERM)
        kills = []

        class FakeOs:
            getpid = staticmethod(os.getpid)

            @staticmethod
            def kill(pid, sig):
                kills.append((pid, sig))

        real_os = mod.os
        mod.os = FakeOs
        try:
            mod._handle_stop_signal(signal.SIGTERM, None)
            self.assertTrue(mod.STOP_EVENT.is_set())
            self.assertEqual(kills, [])
            mod._handle_stop_signal(signal.SIGTERM, None)     # 二次信号：恢复默认处置并当场退出
            self.assertEqual(kills, [(os.getpid(), signal.SIGTERM)])
            self.assertIs(signal.getsignal(signal.SIGTERM), signal.SIG_DFL)
        finally:
            mod.os = real_os
            signal.signal(signal.SIGTERM, orig_term)

    def test_stream_loop_stops_on_stop_event(self):
        mod.STOP_EVENT.set()
        connects = []
        real_connect = mod.connect_okx_ws
        mod.connect_okx_ws = lambda *a, **k: connects.append(1)
        try:
            engine = types.SimpleNamespace(logger=mod.Logger(os.devnull))
            mod.stream_okx(engine, ["TEST-USDT"])            # 应立即返回，不再发起连接
        finally:
            mod.connect_okx_ws = real_connect
        self.assertEqual(connects, [])

    def test_watchdog_restarts_dead_market_thread(self):
        started = []
        mod.stream_okx = lambda engine, ids: started.append(ids)
        engine = types.SimpleNamespace(logger=mod.Logger(os.devnull))
        dead = self._dead_thread()
        t2, last, restarts = mod.watch_market_thread(engine, ["A-USDT"], dead, 0.0, 0, now=1000.0)
        self.assertEqual(restarts, 1)
        self.assertEqual(last, 1000.0)
        self.assertIsNot(t2, dead)
        self.assertEqual(mod.WS_HEALTH["last_error"], "行情线程意外退出")
        t2.join(2)

    def test_watchdog_honors_min_restart_interval(self):
        mod.stream_okx = lambda engine, ids: None
        engine = types.SimpleNamespace(logger=mod.Logger(os.devnull))
        dead = self._dead_thread()
        t2, last, restarts = mod.watch_market_thread(
            engine, ["A-USDT"], dead, 1000.0 - 5.0, 3, now=1000.0)
        self.assertIs(t2, dead)
        self.assertEqual(restarts, 3)                        # 退避窗口内不重复拉起

    def test_watchdog_leaves_alive_thread_alone(self):
        mod.stream_okx = lambda engine, ids: time.sleep(0.5)
        engine = types.SimpleNamespace(logger=mod.Logger(os.devnull))
        alive = mod.start_market_thread(engine, ["A-USDT"])
        t2, last, restarts = mod.watch_market_thread(engine, ["A-USDT"], alive, 0.0, 7, now=1000.0)
        self.assertIs(t2, alive)
        self.assertEqual(restarts, 7)
        self.assertEqual(last, 0.0)
        alive.join(2)


class TestStaleTmpSweep(TempFilesMixin, unittest.TestCase):
    """R8：启动时清理上次运行被 kill 留下的原子写临时文件"""

    def test_sweep_removes_only_stale_tmp(self):
        old = mod._csv_tmp_path(mod.POSITIONS_FILE)
        fresh = mod._csv_tmp_path(mod.HIGH_LOW_FILE)
        other = mod.POSITIONS_FILE + ".bak"
        for p in (old, fresh, other):
            with open(p, "w") as f:
                f.write("x")
        past = time.time() - 3600
        os.utime(old, (past, past))
        removed = mod.sweep_stale_tmp_files()
        self.assertIn(os.path.basename(old), removed)
        self.assertFalse(os.path.exists(old))
        self.assertTrue(os.path.exists(fresh))     # 太新：可能有人在写，不得删
        self.assertTrue(os.path.exists(other))     # 非 tmp 文件一律不碰


class TestConfigDirtyRow(TempFilesMixin, PoolDirMixin, unittest.TestCase):
    """R9：一行脏配置不能掀掉整个程序（否则已有持仓无人盯守）。
    v5.13 股票池格式：sheet2 A=币对、C=同层最低→ref_low、D=同层最高→ref_high"""

    def _load(self, coins):
        # coins: [(inst_id, ref_low, ref_high)]
        self._patch_pool_dir()
        make_pool_workbook(os.path.join(self.pool_dir, "cfg.xlsx"),
                           [(i, lo, hi, "8-16") for i, lo, hi in coins])
        orig_throttle = mod.push_notify_throttled
        mod.push_notify_throttled = lambda *a, **k: None
        try:
            return mod.load_config()
        finally:
            mod.push_notify_throttled = orig_throttle
            self._restore_pool_dir()

    def test_one_dirty_row_does_not_break_others(self):
        config, _ = self._load([
            ("AAA-USDT", 0.5, 2.0),
            ("BAD-USDT", 0.5, "2元"),      # 文本单价
            ("EMPTY-USDT", 0.5, None),     # 空单元格
            ("SWAP-USDT", 2.0, 0.5),       # 高低倒挂
            ("ZERO-USDT", 0, 2.0),         # 低价为 0
        ])
        self.assertEqual(sorted(config), ["AAA-USDT"])

    def test_inst_id_normalized(self):
        config, _ = self._load([("  AAA-USDT  ", 0.5, 2.0)])
        self.assertIn("AAA-USDT", config)

    def test_all_rows_dirty_returns_empty_without_raising(self):
        config, _ = self._load([("BAD-USDT", "x", "y")])
        self.assertEqual(config, {})

    def test_startup_decision(self):
        # 无配置 + 无持仓 → 退出；无配置 + 有持仓 → 必须起来只管卖
        self.assertTrue(mod.config_startup_ok({"AAA-USDT": {}}, []))
        self.assertTrue(mod.config_startup_ok({}, ["AAA-USDT"]))
        self.assertFalse(mod.config_startup_ok({}, []))


class TestLayerFilter(TempFilesMixin, PoolDirMixin, unittest.TestCase):
    """v5.13：FILTER_LAYER 按 sheet1"层计算"区间列（如 8-16，两位数字=本层最高/最低的价差倍数）筛选股票池"""

    def setUp(self):
        super().setUp()
        self._patch_pool_dir()
        orig_throttle = mod.push_notify_throttled
        mod.push_notify_throttled = lambda *a, **k: None
        self.addCleanup(lambda: setattr(mod, "push_notify_throttled", orig_throttle))
        self.addCleanup(self._restore_pool_dir)
        make_pool_workbook(os.path.join(self.pool_dir, "cfg.xlsx"), [
            ("KAT-USDT", 0.0038, 0.0077, "4-8"),
            ("FLR-USDT", 0.0056, 0.0113, "8-16"),
            ("NFT-USDT", 1.98e-07, 3.97e-07, "16-32"),
            ("LAT-USDT", 0.00030, 0.00061, "8192-16384"),
        ])

    def _load_with(self, raw):
        mod.LAYER_FILTERS = mod._parse_layer_filters(raw)
        config, _ = mod.load_config()
        return sorted(config)

    def test_exact_layer(self):
        self.assertEqual(self._load_with("8-16"), ["FLR-USDT"])

    def test_above_layer(self):
        # "8-16及以上"指上一层起：16-32、32-64……（不含 8-16 本身）
        self.assertEqual(self._load_with("8-16及以上"), ["LAT-USDT", "NFT-USDT"])

    def test_below_layer(self):
        # "8-16以下"指下一层起：4-8、2-4……（不含 8-16 本身）
        self.assertEqual(self._load_with("8-16以下"), ["KAT-USDT"])

    def test_multi_conditions_union(self):
        self.assertEqual(self._load_with("4-8,16-32"), ["KAT-USDT", "NFT-USDT"])

    def test_empty_means_no_filter(self):
        self.assertEqual(self._load_with(""), ["FLR-USDT", "KAT-USDT", "LAT-USDT", "NFT-USDT"])

    def test_label_missing_filtered_out(self):
        # 币对在 sheet1 无区间标签、历史最高/最低列又缺失（无法兜底估层）→ 有过滤条件时不入池
        import openpyxl
        path = os.path.join(self.pool_dir, "cfg.xlsx")
        wb = openpyxl.load_workbook(path)
        ws1 = wb["层计算"]
        ws1.cell(row=2, column=3).value = None   # KAT 的区间标签
        ws1.cell(row=2, column=4).value = None   # 历史最低（兜底估层用）
        ws1.cell(row=2, column=6).value = None   # 历史最高
        wb.save(path)
        wb.close()
        self.assertEqual(self._load_with("8-16"), ["FLR-USDT"])

    def test_bad_filter_value_ignored(self):
        # 非法条件字符串解析为空 → 等效不过滤，并点名提示
        self.assertEqual(mod._parse_layer_filters("八-十六"), [])
        self.assertEqual(sorted(self._load_with("八-十六")),
                         ["FLR-USDT", "KAT-USDT", "LAT-USDT", "NFT-USDT"])

    def test_band_fallback_from_history_high_low(self):
        # 区间标签缺失但历史最高/最低在：按 2 的幂层兜底（KAT 0.0077/0.0038≈2.03 倍 → 2-4 层）
        import openpyxl
        path = os.path.join(self.pool_dir, "cfg.xlsx")
        wb = openpyxl.load_workbook(path)
        wb["层计算"].cell(row=2, column=3).value = None
        wb.save(path)
        wb.close()
        self.assertEqual(self._load_with("2-4"), ["KAT-USDT"]) 


class TestBuyQueueStartupBatch(TempFilesMixin, unittest.TestCase):
    """R10：启动期买入队列同步必须是"一次读表 + 一次写表"，不是每币对各重写一遍整表"""

    def _count_writes(self):
        counter = {"n": 0, "rows": None}
        real = mod.save_buy_queue

        def spy(rows):
            counter["n"] += 1
            counter["rows"] = [dict(r) for r in rows]
            return real(rows)

        mod.save_buy_queue = spy
        return counter, real

    def tearDown_write(self, real):
        mod.save_buy_queue = real

    def test_symbol_init_does_not_touch_queue_table(self):
        # 原来 SymbolState.__init__ 与 MarketEngine.add_state 各调一次 _sync_buy_queue：
        # N 个币对 = 2N 次整表读写（还各自 fsync），启动时间随币对数平方增长
        counter, real = self._count_writes()
        try:
            for i in range(8):
                mod.SymbolState(f"S{i}-USDT", 2.0, 0.5, mod.Logger(os.devnull),
                                {"buy_ratio": 0.01, "cooldown_days": 7})
            self.assertEqual(counter["n"], 0)
        finally:
            self.tearDown_write(real)

    def test_sync_all_writes_once_and_is_correct(self):
        engine = mod.MarketEngine(mod.Logger(os.devnull))
        states = {}
        for i in range(12):
            states[f"S{i}-USDT"] = make_state(inst_id=f"S{i}-USDT")
        states["S0-USDT"].has_position = True
        states["S1-USDT"].observe_list = {"S1-USDT"}
        states["S2-USDT"].blacklist = {"S2-USDT"}
        states["S3-USDT"].out_of_range = True
        states["S4-USDT"].position_only = True
        states["S6-USDT"].cooldown = {
            "S6-USDT": (datetime.now() + timedelta(days=1)).strftime("%Y-%m-%d %H:%M:%S")}

        def row(inst_id, status="待买入"):
            return {"inst_id": inst_id, "status": status, "ref_high": "2", "ref_low": "0.5",
                    "add_time": "", "last_check_time": ""}

        # 已有行：S0（转为已持仓）、S2/S3（应被剔除）、S6（转为静默期）
        mod.save_buy_queue([row("S0-USDT"), row("S2-USDT"), row("S3-USDT"), row("S6-USDT")])
        engine.states = states
        counter, real = self._count_writes()
        try:
            self.assertTrue(engine.sync_all_buy_queue())
            self.assertEqual(counter["n"], 1)          # 整表只落一次
        finally:
            self.tearDown_write(real)
        rows = {r["inst_id"]: r["status"] for r in mod.load_buy_queue()}
        self.assertEqual(rows["S0-USDT"], "已持仓_排除")
        self.assertEqual(rows["S6-USDT"], "静默期_排除")
        self.assertEqual(rows["S5-USDT"], "待买入")
        self.assertNotIn("S1-USDT", rows)               # 观察列表：不加新行
        self.assertNotIn("S2-USDT", rows)               # 黑名单：剔除
        self.assertNotIn("S3-USDT", rows)               # 超范围：剔除
        self.assertNotIn("S4-USDT", rows)               # 仅持仓管理：不入队
        self.assertEqual(len(rows), 8)                  # 4 个已有行(2 改状态 2 剔除) + 6 个新进队

    def test_sync_all_skips_write_when_nothing_changed(self):
        engine = mod.MarketEngine(mod.Logger(os.devnull))
        st = make_state(inst_id="S0-USDT")
        ts = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        mod.save_buy_queue([{"inst_id": "S0-USDT", "status": "待买入", "ref_high": str(st.high),
                             "ref_low": str(st.low), "add_time": ts, "last_check_time": ts}])
        engine.states = {"S0-USDT": st}
        counter, real = self._count_writes()
        try:
            self.assertTrue(engine.sync_all_buy_queue())
            self.assertEqual(counter["n"], 0)          # 无改动就不碰磁盘
        finally:
            self.tearDown_write(real)

    def test_single_symbol_sync_still_persists(self):
        # 运行期路径（on_trade / 超范围剔除）仍要能单独落盘
        st = make_state(inst_id="S9-USDT")
        st._sync_buy_queue()
        rows = {r["inst_id"]: r["status"] for r in mod.load_buy_queue()}
        self.assertEqual(rows["S9-USDT"], "待买入")


class TestOutOfRangeDeferredWrite(TempFilesMixin, unittest.TestCase):
    """R11：超范围审计（xlsx 整本重写）挪出行情热路径，改为主线程批量落盘"""

    def setUp(self):
        super().setUp()
        mod._XLSX_OOR_OPS.clear()
        mod._XLSX_SET_CACHE.clear()

    def tearDown(self):
        mod._XLSX_OOR_OPS.clear()
        mod._XLSX_SET_CACHE.clear()
        super().tearDown()

    def _rows(self):
        import openpyxl
        wb = openpyxl.load_workbook(mod.OUT_OF_RANGE_FILE)
        ws = wb.active
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        wb.close()
        return rows

    def test_enqueue_does_not_touch_disk(self):
        mod.add_to_out_of_range("A-USDT", 3.0, 2.0, 0.5)
        mod.mark_recovered_in_out_of_range("A-USDT", 1.0)
        self.assertFalse(os.path.exists(mod.OUT_OF_RANGE_FILE))
        self.assertEqual(len(mod._XLSX_OOR_OPS), 2)

    def test_flush_writes_all_pending_in_one_pass(self):
        for i in range(5):
            mod.add_to_out_of_range(f"S{i}-USDT", 3.0, 2.0, 0.5)
        self.assertEqual(mod.flush_out_of_range_ops(), 5)
        self.assertEqual(len(self._rows()), 6)               # header + 5 行
        self.assertEqual(mod.flush_out_of_range_ops(), 0)    # 队列已空
        self.assertEqual(len(self._rows()), 6)               # 不重复写

    def test_recorded_only_once(self):
        mod.add_to_out_of_range("A-USDT", 3.0, 2.0, 0.5)
        mod.add_to_out_of_range("A-USDT", 4.0, 2.0, 0.5)     # 队列内去重
        self.assertEqual(len(mod._XLSX_OOR_OPS), 1)
        mod.flush_out_of_range_ops()
        mod.add_to_out_of_range("A-USDT", 5.0, 2.0, 0.5)     # 表里已有记录 → 不再加行
        mod.flush_out_of_range_ops()
        self.assertEqual(len(self._rows()), 2)

    def test_recover_rewrites_reason_and_clears_exclusion(self):
        mod.add_to_out_of_range("A-USDT", 3.0, 2.0, 0.5)
        mod.flush_out_of_range_ops()
        self.assertIn("A-USDT", mod.load_out_of_range())
        mod.mark_recovered_in_out_of_range("A-USDT", 1.0)
        mod.flush_out_of_range_ops()
        self.assertNotIn("A-USDT", mod.load_out_of_range())  # 缓存失效，重新解析
        self.assertTrue(str(self._rows()[1][4]).startswith("已恢复"))

    def test_failed_flush_requeues(self):
        def boom(ops):
            raise RuntimeError("模拟写失败")
        real = mod._apply_out_of_range_ops
        mod._apply_out_of_range_ops = boom
        try:
            mod.add_to_out_of_range("A-USDT", 3.0, 2.0, 0.5)
            self.assertEqual(mod.flush_out_of_range_ops(), 0)
            self.assertEqual(len(mod._XLSX_OOR_OPS), 1)      # 事件回队，不静默丢审计
        finally:
            mod._apply_out_of_range_ops = real
        self.assertEqual(mod.flush_out_of_range_ops(), 1)
        self.assertIn("A-USDT", mod.load_out_of_range())
        self.assertFalse(os.path.exists(mod._csv_tmp_path(mod.OUT_OF_RANGE_FILE)))

    def test_layer_switch_replaces_out_of_range(self):
        # v6（设计单 v1.1）：价格越出层界 → 换层（不再剔除、不再写超范围审计表）
        st = make_state("A-USDT")
        st._save_state = lambda force=False: None
        st._push_notify = lambda t: None
        st.update_price(3.0, 1.0)                            # 超出层界上沿 2.0 → 向上换层
        self.assertAlmostEqual(st.levels.low, 1.0)
        self.assertAlmostEqual(st.levels.high, 4.0)
        self.assertEqual(st.v6_next_buy, "next_layer")       # 无持仓向上换层：忽略（不高位接盘）
        self.assertEqual(len(mod._XLSX_OOR_OPS), 0)
        self.assertFalse(os.path.exists(mod.OUT_OF_RANGE_FILE))


class TestWsStallReconnect(unittest.TestCase):
    """R12：TCP 没断但链路静默（收不到任何行情/pong）时必须主动重连，不能只发 ping 干等"""

    def test_action_thresholds(self):
        self.assertEqual(mod.ws_stall_action(mod.PING_INTERVAL - 1), "")
        self.assertEqual(mod.ws_stall_action(mod.PING_INTERVAL), "ping")
        self.assertEqual(mod.ws_stall_action(mod.WS_STALL_TIMEOUT_SEC - 1), "ping")
        self.assertEqual(mod.ws_stall_action(mod.WS_STALL_TIMEOUT_SEC), "stall")

    def test_silent_link_breaks_to_reconnect(self):
        closed = {"n": 0}

        class SilentConn:
            def recv_message(self):
                raise socket.timeout()

            def send_text(self, msg):
                pass

            def close(self):
                closed["n"] += 1
                mod.STOP_EVENT.set()      # 让外层循环收尾，测试只跑一轮

        real = {"connect": mod.connect_okx_ws, "sub": mod.subscribe_okx_trades,
                "throttle": mod.push_notify_throttled, "stall": mod.WS_STALL_TIMEOUT_SEC,
                "health": dict(mod.WS_HEALTH), "stop": mod.STOP_EVENT.is_set()}
        mod.connect_okx_ws = lambda *a, **k: SilentConn()
        mod.subscribe_okx_trades = lambda conn, ids: None
        mod.push_notify_throttled = lambda *a, **k: None
        mod.WS_STALL_TIMEOUT_SEC = 0.0
        mod.STOP_EVENT.clear()
        seen = {}
        try:
            engine = types.SimpleNamespace(logger=mod.Logger(os.devnull))
            mod.stream_okx(engine, ["TEST-USDT"])
            seen = {"last_error": mod.WS_HEALTH["last_error"],
                    "connected": mod.WS_HEALTH["connected"],
                    "stall": mod.WS_HEALTH.get("stall_reconnects", 0)}
        finally:
            mod.connect_okx_ws = real["connect"]
            mod.subscribe_okx_trades = real["sub"]
            mod.push_notify_throttled = real["throttle"]
            mod.WS_STALL_TIMEOUT_SEC = real["stall"]
            mod.WS_HEALTH.clear()
            mod.WS_HEALTH.update(real["health"])
            mod.STOP_EVENT.clear()
            if real["stop"]:
                mod.STOP_EVENT.set()
        self.assertEqual(closed["n"], 1)                      # 静默链路被关闭
        self.assertIn("静默", seen["last_error"])              # 心跳里留下静默原因
        self.assertFalse(seen["connected"])
        self.assertEqual(seen["stall"], 1)


class TestEnvNum(unittest.TestCase):
    """R15：.env 里一个写错的数值不该让程序起不来（启动即 traceback = 所有持仓无人盯守）"""

    NAME = "OKX_TEST_ENV_NUM"

    def tearDown(self):
        os.environ.pop(self.NAME, None)

    def test_valid_and_blank(self):
        os.environ[self.NAME] = " 12.5 "
        self.assertEqual(mod.env_num(self.NAME, 3.0), 12.5)
        os.environ[self.NAME] = ""
        self.assertEqual(mod.env_num(self.NAME, 3.0), 3.0)
        os.environ.pop(self.NAME)
        self.assertEqual(mod.env_num(self.NAME, 3.0), 3.0)

    def test_bad_value_falls_back_without_raising(self):
        os.environ[self.NAME] = "abc"
        self.assertEqual(mod.env_num(self.NAME, 0.005), 0.005)

    def test_int_cast(self):
        os.environ[self.NAME] = "7"
        self.assertEqual(mod.env_num(self.NAME, 3, int), 7)
        os.environ[self.NAME] = "7.5"          # int("7.5") 会抛 → 回退默认，不能让 main 崩
        self.assertEqual(mod.env_num(self.NAME, 3, int), 3)


class TestRoundDown(unittest.TestCase):
    """R15：按 lot_sz 向下取整改用 Decimal，修掉 float 除法少算一个步长"""

    rd = staticmethod(mod.OkxTradeClient._round_down)

    def test_exact_multiple_not_lost_to_float(self):
        # 0.3/0.1 在二进制浮点里是 2.9999999999999996，旧实现会算成 0.2 个 lot
        self.assertAlmostEqual(self.rd(0.3, 0.1), 0.3, places=12)
        self.assertAlmostEqual(self.rd("0.3", "0.1"), 0.3, places=12)

    def test_floor_still_floors(self):
        self.assertAlmostEqual(self.rd(0.37, "0.1"), 0.3, places=12)
        self.assertAlmostEqual(self.rd(123.456, "1"), 123.0, places=12)
        self.assertAlmostEqual(self.rd(0.09, "0.0001"), 0.09, places=12)

    def test_scientific_notation_step(self):
        # 旧实现按 str(step) 数小数位：1e-08 没有小数点 → decimals=0 → 数量被取整成整数
        self.assertAlmostEqual(self.rd(1.23456, 1e-08), 1.23456, places=12)

    def test_bad_step_returns_value(self):
        for step in ("0", -1, "", None, "x"):
            self.assertEqual(self.rd(1.234, step), 1.234)


class TestRangeRecoverBand(unittest.TestCase):
    """R15：区间比滞回带还窄时，恢复带不能塌成空集（否则一次越界永久剔除）"""

    def test_wide_range_keeps_configured_pct(self):
        lo, hi, pct = mod.range_recover_band(0.5, 2.0)
        self.assertAlmostEqual(pct, 0.005)
        self.assertAlmostEqual(lo, 0.5025)
        self.assertAlmostEqual(hi, 1.99)

    def test_narrow_band_is_not_empty_and_inside(self):
        low, high = 1.0, 1.005                      # high/low=1.005 < (1+p)/(1-p)≈1.01005
        old_lo, old_hi = low * 1.005, high * 0.995  # 旧公式：下界已越过上界
        self.assertGreater(old_lo, old_hi)
        lo, hi, pct = mod.range_recover_band(low, high)
        self.assertLess(lo, hi)
        self.assertGreater(lo, low)
        self.assertLess(hi, high)
        self.assertLessEqual(pct, 0.005)
        mid = (low + high) / 2
        self.assertTrue(lo <= mid <= hi)            # 回到区间中部即可恢复

    def test_degenerate_input_returns_raw_range(self):
        self.assertEqual(mod.range_recover_band(0, 5)[:2], (0, 5))
        self.assertEqual(mod.range_recover_band(5, 5)[:2], (5, 5))
        self.assertEqual(mod.range_recover_band(1.0, 2.0, 0)[:2], (1.0, 2.0))


class TestBuyBalanceGate(TempFilesMixin, unittest.TestCase):
    """R15：余额闸门跟随 MIN_BUY_AMOUNT，且冷却期内不再每 tick 打一次余额 REST"""

    class GateClient:
        def __init__(self, balance):
            self.balance = balance
            self.calls = 0
            self.buy_calls = 0

        def get_balance(self, ccy="USDT"):
            self.calls += 1
            return self.balance

        def get_instrument(self, inst_id):
            return {"minSz": "0.0001", "lotSz": "0.0001", "baseCcy": "TEST"}

        def market_buy_quote(self, inst_id, amount, px=None):
            # 返回"未成交"让流程安全落地：本用例只关心余额闸门有没有放行
            self.buy_calls += 1
            return {"not_filled": True, "state": "test"}

    def _gate_state(self, balance, min_buy=100.0):
        st = make_state("TEST-USDT")
        st.client = self.GateClient(balance)
        st.min_buy_amount = min_buy
        st._push_notify = lambda t: None
        st._save_state = lambda force=False: None
        st._remove_from_buy_queue = lambda: None
        st._add_to_buy_queue = lambda s: None
        orig = mod._okx_client
        mod._okx_client = st.client
        self.addCleanup(setattr, mod, "_okx_client", orig)
        return st

    def _try(self, st):
        return st._execute_buy(1.0, "测试买入", "pullback2_down", "pullback_up", "回撤")

    def test_low_balance_sets_cooldown(self):
        st = self._gate_state(50.0)
        self.assertIsNone(self._try(st))
        self.assertGreater(st._balance_insufficient_until, time.time())

    def test_cooldown_skips_rest_query(self):
        st = self._gate_state(50.0)
        self._try(st)
        self.assertEqual(st.client.calls, 1)
        self._try(st)
        self.assertEqual(st.client.calls, 1)          # 冷却期内一个 tick 都不该再问

    def test_gate_follows_min_buy_amount(self):
        # 旧实现写死 110：min_buy_amount=500、余额 300 时会一路走到下单再被 OKX 拒
        st = self._gate_state(300.0, min_buy=500.0)
        self.assertIsNone(self._try(st))
        self.assertEqual(st.client.buy_calls, 0)
        self.assertGreater(st._balance_insufficient_until, time.time())

    def test_default_gate_matches_legacy_110(self):
        st = self._gate_state(120.0)                  # 120 ≥ 100×1.1 → 放行（与 v5.11 一致）
        self._try(st)
        self.assertEqual(st.client.calls, 1)
        self.assertEqual(st.client.buy_calls, 1)
        self.assertEqual(st._balance_insufficient_until, 0.0)


class TestLadderSpec(unittest.TestCase):
    """R4 规格守卫：价位表按用户 2026-10-12 的声明逐条锁定，不再靠"顺序看着对"

    声明原文（"A+B" 指 A、B 的中间值）：
      最高 > 高位中间 > 中间 > 低位中间 > 最低，族内 上 > 中 > 下
      高位中间=最高+中间   中间=最高+最低   低位中间=中间+最低
      回撤上=最高+回撤中   回撤中=最高+高位中间   回撤下=回撤中+高位中间
      回调上=高位中间+回调中  回调中=高位中间+中间  回调下=回调中+中间
      过渡上=中间+过渡中     过渡中=中间+低位中间   过渡下=过渡中+低位中间
      极限上=低位中间+极限中  极限中=低位中间+最低   极限下=极限中+最低
    """

    RANGES = [(0.0981281, 0.0490641), (2.0, 0.5), (100.0, 50.0), (7.3, 1.1), (1e-7, 5e-8)]

    @staticmethod
    def declared(high, low):
        M = lambda a, b: (a + b) / 2
        d = {"中间": M(high, low)}
        d["高位中间"] = M(high, d["中间"])
        d["低位中间"] = M(d["中间"], low)
        for fam, top, bottom in (("回撤", high, "高位中间"), ("回调", "高位中间", "中间"),
                                 ("过渡", "中间", "低位中间"), ("极限", "低位中间", low)):
            d[f"{fam}中"] = M(top if isinstance(top, float) else d[top],
                              bottom if isinstance(bottom, float) else d[bottom])
            d[f"{fam}上"] = M(top if isinstance(top, float) else d[top], d[f"{fam}中"])
            d[f"{fam}下"] = M(d[f"{fam}中"], bottom if isinstance(bottom, float) else d[bottom])
        return d

    CODE_ATTR = {
        "回撤上": "pullback_up", "回撤中": "pullback1_mid", "回撤下": "pullback2_down",
        "高位中间": "high_mid",
        "回调上": "callback_up", "回调中": "callback1_mid", "回调下": "callback2_down",
        "中间": "mid",
        "过渡上": "transition_up", "过渡中": "transition1_mid", "过渡下": "transition2_down",
        "低位中间": "low_mid",
        "极限上": "limit_up", "极限中": "limit1_mid", "极限下": "limit2_down",
    }

    def test_every_declared_level(self):
        for high, low in self.RANGES:
            lv = mod.PriceLevels(high, low)
            d = self.declared(high, low)
            for name, attr in self.CODE_ATTR.items():
                self.assertAlmostEqual(getattr(lv, attr), d[name],
                                       msg=f"{high}/{low} 的 {name}", delta=abs(high) * 1e-12)

    def test_anchor_ordering(self):
        for high, low in self.RANGES:
            lv = mod.PriceLevels(high, low)
            chain = [high, lv.high_mid, lv.mid, lv.low_mid, low]
            names = ["最高", "高位中间", "中间", "低位中间", "最低"]
            for i in range(4):
                self.assertGreater(chain[i], chain[i + 1], f"{names[i]} 应大于 {names[i+1]}")

    def test_family_ordering_up_mid_down(self):
        for high, low in self.RANGES:
            lv = mod.PriceLevels(high, low)
            for fam in ("pullback", "callback", "transition", "limit"):
                up, mid, down = (getattr(lv, f"{fam}_up"), getattr(lv, f"{fam}1_mid"),
                                 getattr(lv, f"{fam}2_down"))
                self.assertGreater(up, mid, f"{fam} 上应大于中")
                self.assertGreater(mid, down, f"{fam} 中应大于下")

    def test_buy_and_sell_levels_are_declared_names(self):
        # 引擎只用 *_下 作买点、*_上 作卖点（回调下 = 回调买点，回调上 = 回调卖点）
        lv = mod.PriceLevels(0.0981281, 0.0490641)
        self.assertEqual(set(lv.buy_levels), {"pullback2_down", "callback2_down",
                                              "transition2_down", "limit2_down"})
        self.assertEqual(set(lv.sell_levels), {"pullback_up", "callback_up",
                                               "transition_up", "limit_up"})
        for name, (sell_key, cn) in lv.pair_map.items():
            self.assertEqual(lv.sell_levels[sell_key], getattr(lv, self.CODE_ATTR[f"{cn}上"]), f"{cn}卖点")
            self.assertEqual(lv.buy_levels[name], getattr(lv, self.CODE_ATTR[f"{cn}下"]), f"{cn}买点")


    def test_no_undeclared_levels(self):
        # "丢弃旧版"（用户 2026-10-12 20:35）：权威表以外的价位属性不得复活
        lv = mod.PriceLevels(2.0, 0.5)
        have = {a for a, v in vars(lv).items() if isinstance(v, float)}
        want = set(self.CODE_ATTR.values()) | {"high", "low"}
        self.assertEqual(have, want, f"多出 {have - want}，缺少 {want - have}")

    def test_full_declared_ordering_chain(self):
        """用户声明的 17 点全序（2026-10-12 20:30 版，含最高/最低两端）"""
        chain = [
            ("最高", "high"),
            ("回撤上", "pullback_up"), ("回撤中", "pullback1_mid"), ("回撤下", "pullback2_down"),
            ("高位中间", "high_mid"),
            ("回调上", "callback_up"), ("回调中", "callback1_mid"), ("回调下", "callback2_down"),
            ("中间", "mid"),
            ("过渡上", "transition_up"), ("过渡中", "transition1_mid"), ("过渡下", "transition2_down"),
            ("低位中间", "low_mid"),
            ("极限上", "limit_up"), ("极限中", "limit1_mid"), ("极限下", "limit2_down"),
            ("最低", "low"),
        ]
        for high, low in self.RANGES:
            lv = mod.PriceLevels(high, low)
            values = [(name, getattr(lv, attr)) for name, attr in chain]
            for i in range(len(values) - 1):
                hi_name, hi_v = values[i]
                lo_name, lo_v = values[i + 1]
                self.assertGreater(hi_v, lo_v,
                                   f"{high}/{low}: 声明 {hi_name} > {lo_name}，实际 {hi_v} ≤ {lo_v}")

    def test_declared_chain_is_evenly_spaced(self):
        """每个价位都是相邻锚点的中点 ⇒ 15 个价位把区间十六等分，间距应处处相等"""
        for high, low in self.RANGES:
            lv = mod.PriceLevels(high, low)
            names = ["high", "pullback_up", "pullback1_mid", "pullback2_down", "high_mid",
                     "callback_up", "callback1_mid", "callback2_down", "mid",
                     "transition_up", "transition1_mid", "transition2_down", "low_mid",
                     "limit_up", "limit1_mid", "limit2_down", "low"]
            vals = [getattr(lv, n) for n in names]
            steps = [vals[i] - vals[i + 1] for i in range(len(vals) - 1)]
            expect = (high - low) / 16
            for i, s in enumerate(steps):
                self.assertAlmostEqual(s, expect, delta=abs(high) * 1e-12,
                                       msg=f"{names[i]} → {names[i+1]} 间距 {s} ≠ {expect}")


if __name__ == "__main__":
    unittest.main(verbosity=2)
